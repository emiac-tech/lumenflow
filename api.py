from __future__ import annotations

import json
import os
import re
from io import BytesIO
from decimal import Decimal, InvalidOperation
import shutil
import uuid
from contextlib import asynccontextmanager, contextmanager
from pathlib import Path
from typing import List, Literal, Optional
from urllib.parse import unquote, urlparse
from xml.sax.saxutils import escape

import requests
import uvicorn
from dotenv import load_dotenv
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from fastapi.staticfiles import StaticFiles
from sqlalchemy import create_engine, text
from sqlalchemy.orm import DeclarativeBase, Mapped, Session, mapped_column, sessionmaker

from extract_images import process_luminaire_document
from webpage_extract import (
    MAX_PAGE_BYTES,
    MAX_PDF_BYTES,
    decode_html,
    looks_like_pdf,
    process_luminaire_webpage,
    safe_get,
)
import asyncio
from concurrent.futures import ThreadPoolExecutor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A3, A4, LETTER, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.utils import ImageReader, simpleSplit
from reportlab.pdfgen import canvas
from reportlab.platypus import Image as RLImage
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from datetime import datetime, timedelta
from sqlalchemy import JSON, DateTime, ForeignKey, String, func
from sqlalchemy.orm import relationship

from passlib.context import CryptContext
from jose import JWTError, jwt
from fastapi.security import OAuth2PasswordBearer, OAuth2PasswordRequestForm
from fastapi import Depends
from pydantic import BaseModel
from google.oauth2 import id_token
from google.auth.transport import requests as google_requests
from default_spec_template import (
    DEFAULT_TEMPLATE_KEY,
    DEFAULT_TEMPLATE_NAME,
    DEFAULT_TEMPLATE_VERSION,
    PAGE_HEIGHT as DEFAULT_TEMPLATE_HEIGHT,
    PAGE_WIDTH as DEFAULT_TEMPLATE_WIDTH,
    build_default_specification_template_design,
    COMPACT_TEMPLATE_KEY,
    COMPACT_TEMPLATE_NAME,
    TEAL as DEFAULT_TEMPLATE_ACCENT,
    build_compact_specification_template_design,
)


BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")
DATA_DIR = BASE_DIR / "data"
DOCUMENTS_DIR = DATA_DIR / "documents"
PRODUCT_IMAGES_DIR = DATA_DIR / "images" / "products"
DIAGRAM_IMAGES_DIR = DATA_DIR / "images" / "diagrams"
POLAR_IMAGES_DIR = DATA_DIR / "images" / "polar"
DESIGNER_IMAGES_DIR = DATA_DIR / "images" / "designer"
EXPORTS_DIR = DATA_DIR / "exports"
BRANDING_DIR = DATA_DIR / "branding"
SCHEDULE_BRAND_LOGO_SOURCE_PATH = BRANDING_DIR / "afterglow-logo.jpeg"
SCHEDULE_BRAND_LOGO_EXPORT_PATH = BRANDING_DIR / "afterglow-logo-prepared.png"

for folder in [DATA_DIR, DOCUMENTS_DIR, PRODUCT_IMAGES_DIR, DIAGRAM_IMAGES_DIR, POLAR_IMAGES_DIR, DESIGNER_IMAGES_DIR, EXPORTS_DIR, BRANDING_DIR]:
    folder.mkdir(parents=True, exist_ok=True)


SCHEDULE_EXPORT_LABELS = {
    "type": "Type",
    "photo": "Photo",
    "diagram": "Diagram",
    "polar": "Polar Curve",
    "manufacturer": "Manufacturer",
    "product": "Product",
    "specCode": "Spec Code",
    "description": "Description",
    "colour": "Colour",
    "cct": "CCT",
    "cri": "CRI",
    "unitWatts": "Unit Watts",
    "dimming": "Dimming",
    "quantity": "Qty",
    "ipRating": "IP Rating",
    "totalWatts": "Total Watts",
    "document": "Document",
    "templatePdf": "Template PDF",
}

IMPORT_METADATA_ALIASES = {
    "type": ("type", "typecode", "fittingtype", "fitting", "luminairetype"),
    "manufacturer": ("manufacturer", "manufacturername", "brand", "brandname", "make", "vendor", "company"),
    "product": ("product", "productname", "model", "modelname", "productmodel"),
    "spec_code": ("speccode", "specificationcode", "productcode", "modelcode", "referencecode"),
    "description": ("description", "desc", "productdescription"),
    "colour": ("colour", "color", "finish"),
    "cct": ("cct", "ledcct", "ledcricct", "cricct", "colourtemperature", "colortemperature", "colourtemperatures", "colortemperatures", "cctoptions", "cctcolortemperatureoptions", "cctcolourtemperatureoptions", "colortemperatureoptions", "colourtemperatureoptions", "lightcolour", "lightcolor"),
    "cri": ("cri", "colourrenderingindex", "colorrenderingindex"),
    "unit_watts": ("unitwatts", "unitwatt", "maxwattage", "maximumwattage", "wattage", "watts", "power", "systempower", "inputpower", "ratedpower", "circuitpower"),
    "dimming": ("dimming", "dimmingtype", "dimmingoptions", "dimmingoption", "dimmingprotocols", "dimmingprotocol", "optionaldimming", "dimmable", "control", "controloptions", "controlgear"),
    "quantity": ("quantity", "qty", "count"),
    "ip_rating": ("iprating", "ip", "ingressprotection", "degreeofprotection", "protectiverating", "protectionrating", "technicalprotectionratings", "technicalandprotectionratings"),
    "total_watts": ("totalwatts", "totalwattage", "totalpower"),
    "template_pdf": ("templatepdf", "template", "templatefile", "templatepdfformat"),
}

LUMINAIRE_REFERENCE_COLUMN_ALIASES = {
    "luminaireref",
    "luminairereference",
}

DEFAULT_SPECIFICATION_COLUMNS = [
    {
        "name": "Location",
        "aliases": ("Luminaire Location", "Room Location", "Area"),
        "description": "Project room, area, floor, or installation location only when explicitly stated in the uploaded document.",
    },
    {
        "name": "Light Source",
        "aliases": ("Lamp Type", "LED Type", "Lamp Code"),
        "description": "Light source or lamp technology, such as LED, COB LED, fluorescent, or replaceable lamp.",
    },
    {
        "name": "MacAdam / SDCM",
        "aliases": ("MacAdam", "SDCM", "MacAdam Step"),
        "description": "LED colour consistency value expressed as MacAdam steps or SDCM.",
    },
    {
        "name": "Beam / Distribution",
        "aliases": ("Beam", "Beam Angle", "Distribution", "Optic"),
        "description": "Beam angle, optic name, or light distribution, including narrow/wide/flood/wall-wash descriptions.",
    },
    {
        "name": "Delivered Lumens",
        "aliases": ("Lumens", "Luminous Flux", "System Lumens", "Light Output"),
        "description": "Delivered or system luminous flux in lumens; prefer luminaire/system lumens over raw source lumens.",
    },
    {
        "name": "LED Lifetime",
        "aliases": ("Lifetime", "Rated Life", "Life Time LED"),
        "description": "LED lifetime including hours and maintenance/failure notation such as L80, L90, B10, and ambient temperature.",
    },
    {
        "name": "Input Voltage",
        "aliases": ("Voltage", "Supply Voltage", "Voltage [Vin]"),
        "description": "Required input or supply voltage, including AC/DC and frequency when provided.",
    },
    {
        "name": "Emergency Provision",
        "aliases": ("Emergency", "Emergency Mode"),
        "description": "Emergency-lighting provision, emergency output, battery duration, or compatibility; leave empty when absent.",
    },
    {
        "name": "Dimensions",
        "aliases": ("Dimension", "Size", "Overall Dimensions"),
        "description": "Overall product dimensions from text or the dimension drawing, with units and orientation where available.",
    },
    {
        "name": "Material",
        "aliases": ("Housing Material", "Body Material", "Construction"),
        "description": "Primary luminaire body, housing, cover, diffuser, or optic materials stated in the document.",
    },
    {
        "name": "Dimming / Control",
        "aliases": ("Control", "Control Protocol", "Dimming Protocol"),
        "description": "Dimming or control protocol such as DALI, DALI-2, DMX/RDM, 0-10V, Casambi, phase dimming, or switchable.",
    },
    {
        "name": "IK Rating",
        "aliases": ("IK", "Impact Rating"),
        "description": "Mechanical impact protection rating in IK format, for example IK08.",
    },
    {
        "name": "Weight",
        "aliases": ("Weight (Kg)", "Product Weight", "Net Weight"),
        "description": "Product or luminaire weight with unit; do not return packaging or shipping weight.",
    },
    {
        "name": "Operating Temperature",
        "aliases": ("Operating Temp.", "Ambient Temperature", "Ta Range"),
        "description": "Permitted ambient or operating temperature range with units.",
    },
    {
        "name": "Required Accessories",
        "aliases": ("Accessories", "Required Accessory", "Installation Accessories"),
        "description": "Accessories, bases, drivers, power supplies, mounting parts, or control gear required for installation or operation.",
    },
    {
        "name": "General Notes",
        "aliases": ("Notes", "Remarks", "Compatibility Notes"),
        "description": "Important product notes, compatibility statements, limitations, or installation remarks not captured by another field.",
    },
]


DATABASE_URL = os.getenv("DATABASE_URL", f"sqlite:///{BASE_DIR / 'luminaire.db'}")
FREEMIUM_DEFAULT_CREDITS = max(int(os.getenv("FREEMIUM_DEFAULT_CREDITS", "5")), 0)
_is_sqlite = DATABASE_URL.startswith("sqlite")
engine = create_engine(
    DATABASE_URL,
    connect_args={"check_same_thread": False} if _is_sqlite else {},
    pool_pre_ping=True,           # test connection health before use
    pool_recycle=1800,            # recycle connections after 30 min to avoid SSL EOF from remote DB
)
SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)


class Base(DeclarativeBase):
    pass


class User(Base):
    __tablename__ = "users"
    id: Mapped[int] = mapped_column(primary_key=True, autoincrement=True)
    email: Mapped[str] = mapped_column(unique=True, index=True)
    hashed_password: Mapped[Optional[str]]
    role: Mapped[str] = mapped_column(default="client") # "admin" or "client"
    account_type: Mapped[str] = mapped_column(default="invited") # "admin", "invited", or "freemium"
    is_active: Mapped[bool] = mapped_column(default=True)
    google_sub: Mapped[Optional[str]] = mapped_column(unique=True, index=True)
    tokens: Mapped[int] = mapped_column(default=0)
    invite_code_id: Mapped[Optional[int]] = mapped_column(ForeignKey("invite_codes.id"), index=True)
    created_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)
    name: Mapped[Optional[str]] = mapped_column(String, nullable=True)
    studio_name: Mapped[Optional[str]] = mapped_column(String, nullable=True)

    projects: Mapped[List["Project"]] = relationship(back_populates="owner", cascade="all, delete-orphan")
    invite_code: Mapped[Optional["InviteCode"]] = relationship(back_populates="users")


class InviteCode(Base):
    __tablename__ = "invite_codes"
    id: Mapped[int] = mapped_column(primary_key=True, autoincrement=True)
    code: Mapped[str] = mapped_column(unique=True, index=True)
    token_amount: Mapped[int] = mapped_column(default=1)
    usage_limit: Mapped[int] = mapped_column(default=1)
    is_active: Mapped[bool] = mapped_column(default=True)
    created_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)

    users: Mapped[List["User"]] = relationship(back_populates="invite_code")


class Project(Base):
    __tablename__ = "projects"
    id: Mapped[int] = mapped_column(primary_key=True, autoincrement=True)
    owner_id: Mapped[Optional[int]] = mapped_column(ForeignKey("users.id"), index=True)
    name: Mapped[str]
    created_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)
    
    owner: Mapped[Optional["User"]] = relationship(back_populates="projects")
    luminaires: Mapped[List["Luminaire"]] = relationship(back_populates="project", cascade="all, delete-orphan")
    custom_columns: Mapped[List["CustomColumn"]] = relationship(back_populates="project", cascade="all, delete-orphan")
    templates: Mapped[List["DesignTemplate"]] = relationship(back_populates="project", cascade="all, delete-orphan")
    settings: Mapped[List["AppSetting"]] = relationship(back_populates="project", cascade="all, delete-orphan")


class Luminaire(Base):
    __tablename__ = "luminaires"

    id: Mapped[int] = mapped_column(primary_key=True, autoincrement=True)
    project_id: Mapped[Optional[int]] = mapped_column(ForeignKey("projects.id"), index=True)
    source_type: Mapped[str]
    source_name: Mapped[str]
    source_url: Mapped[Optional[str]]
    document_path: Mapped[str]
    type_code: Mapped[Optional[str]]
    manufacturer: Mapped[Optional[str]]
    product_name: Mapped[Optional[str]]
    spec_code: Mapped[Optional[str]]
    description: Mapped[Optional[str]]
    colour: Mapped[Optional[str]]
    cct: Mapped[Optional[str]]
    cri: Mapped[Optional[str]]
    unit_watts: Mapped[Optional[str]]
    dimming: Mapped[Optional[str]]
    quantity: Mapped[Optional[str]]
    ip_rating: Mapped[Optional[str]]
    total_watts: Mapped[Optional[str]]
    diagram_image_path: Mapped[Optional[str]]
    product_image_path: Mapped[Optional[str]]
    polar_image_path: Mapped[Optional[str]]
    custom_data: Mapped[Optional[str]] = mapped_column(default="{}")
    pdf_template: Mapped[Optional[str]] = mapped_column(default="modern")
    
    project: Mapped[Project] = relationship(back_populates="luminaires")
    templates: Mapped[List["DesignTemplate"]] = relationship(back_populates="luminaire", cascade="all, delete-orphan")
    exports: Mapped[List["GeneratedDocument"]] = relationship(back_populates="luminaire", cascade="all, delete-orphan")

class DesignTemplate(Base):
    __tablename__ = "design_templates"
    id: Mapped[int] = mapped_column(primary_key=True, autoincrement=True)
    project_id: Mapped[Optional[int]] = mapped_column(ForeignKey("projects.id"), index=True)
    luminaire_id: Mapped[Optional[int]] = mapped_column(ForeignKey("luminaires.id"), index=True, nullable=True)
    name: Mapped[str]
    page_size: Mapped[str] = mapped_column(default="A4")
    page_width: Mapped[int] = mapped_column(default=794)
    page_height: Mapped[int] = mapped_column(default=1123)
    design_json: Mapped[dict] = mapped_column(JSON, default=dict)
    created_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)
    updated_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    
    project: Mapped[Project] = relationship(back_populates="templates")
    luminaire: Mapped[Luminaire] = relationship(back_populates="templates")
    exports: Mapped[List["GeneratedDocument"]] = relationship(back_populates="template", cascade="all, delete-orphan")

class GeneratedDocument(Base):
    __tablename__ = "generated_documents"
    id: Mapped[int] = mapped_column(primary_key=True, autoincrement=True)
    luminaire_id: Mapped[int] = mapped_column(ForeignKey("luminaires.id"), index=True)
    template_id: Mapped[int] = mapped_column(ForeignKey("design_templates.id"), index=True)
    title: Mapped[str]
    file_path: Mapped[str]
    created_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)
    luminaire: Mapped[Luminaire] = relationship(back_populates="exports")
    template: Mapped[DesignTemplate] = relationship(back_populates="exports")

class AppSetting(Base):
    __tablename__ = "app_settings"
    id: Mapped[int] = mapped_column(primary_key=True, autoincrement=True)
    project_id: Mapped[Optional[int]] = mapped_column(ForeignKey("projects.id"), index=True)
    key: Mapped[str]
    value: Mapped[Optional[str]]
    updated_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)

    project: Mapped[Project] = relationship(back_populates="settings")

class CustomColumn(Base):
    __tablename__ = "custom_columns"
    id: Mapped[int] = mapped_column(primary_key=True, autoincrement=True)
    project_id: Mapped[Optional[int]] = mapped_column(ForeignKey("projects.id"), index=True)
    name: Mapped[str]
    ai_enabled: Mapped[bool] = mapped_column(default=False)
    ai_description: Mapped[Optional[str]]

    project: Mapped[Project] = relationship(back_populates="custom_columns")


class ScheduleExportRequest(BaseModel):
    luminaire_ids: List[int]
    column_keys: List[str]
    column_labels: dict[str, str] = {}
    export_format: Literal["excel", "pdf"]


@asynccontextmanager
async def lifespan(_: FastAPI):
    Base.metadata.create_all(bind=engine)
    # Automatically add custom_data and pdf_template if older table formulation
    with engine.connect() as conn:
        try:
            conn.execute(text("ALTER TABLE luminaires ADD COLUMN project_id INTEGER REFERENCES projects(id)"))
            conn.commit()
        except Exception:
            conn.rollback()

        try:
            conn.execute(text("ALTER TABLE luminaires ADD COLUMN custom_data TEXT DEFAULT '{}'"))
            conn.commit()
        except Exception:
            conn.rollback()
        try:
            conn.execute(text("ALTER TABLE luminaires ADD COLUMN pdf_template TEXT DEFAULT 'modern'"))
            conn.commit()
        except Exception:
            conn.rollback()

        try:
            conn.execute(text("ALTER TABLE luminaires ADD COLUMN polar_image_path TEXT"))
            conn.commit()
        except Exception:
            conn.rollback()
        
        # New multi-project migration
        try:
            conn.execute(text("ALTER TABLE custom_columns ADD COLUMN project_id INTEGER REFERENCES projects(id)"))
            conn.commit()
        except Exception:
            conn.rollback()

        try:
            conn.execute(text("ALTER TABLE custom_columns ADD COLUMN ai_enabled BOOLEAN DEFAULT FALSE"))
            conn.commit()
        except Exception:
            conn.rollback()

        try:
            conn.execute(text("ALTER TABLE custom_columns ADD COLUMN ai_description TEXT"))
            conn.commit()
        except Exception:
            conn.rollback()

        try:
            conn.execute(text("ALTER TABLE design_templates ADD COLUMN project_id INTEGER REFERENCES projects(id)"))
            conn.commit()
        except Exception:
            conn.rollback()

        try:
            conn.execute(text("ALTER TABLE design_templates ALTER COLUMN luminaire_id DROP NOT NULL"))
            conn.commit()
        except Exception:
            conn.rollback()

        try:
            # Recreate app_settings if it's the old format (key as PK)
            # This is tricky without alembic, but we can try to add columns
            conn.execute(text("ALTER TABLE app_settings ADD COLUMN id INTEGER"))
            conn.execute(text("ALTER TABLE app_settings ADD COLUMN project_id INTEGER REFERENCES projects(id)"))
            conn.commit()
        except Exception:
            conn.rollback()

        # Add owner_id to projects
        try:
            conn.execute(text("ALTER TABLE projects ADD COLUMN owner_id INTEGER REFERENCES users(id)"))
            conn.commit()
        except Exception:
            conn.rollback()

        # Invite-code/account-credit migration
        try:
            conn.execute(text("ALTER TABLE users ADD COLUMN tokens INTEGER DEFAULT 0"))
            conn.commit()
        except Exception:
            conn.rollback()

        try:
            conn.execute(text("ALTER TABLE users ADD COLUMN invite_code_id INTEGER REFERENCES invite_codes(id)"))
            conn.commit()
        except Exception:
            conn.rollback()

        for statement in [
            "ALTER TABLE users ADD COLUMN name VARCHAR",
            "ALTER TABLE users ADD COLUMN studio_name VARCHAR",
            "ALTER TABLE users ADD COLUMN account_type VARCHAR DEFAULT 'invited'",
            "ALTER TABLE users ADD COLUMN is_active BOOLEAN DEFAULT TRUE",
            "ALTER TABLE users ALTER COLUMN account_type SET DEFAULT 'invited'",
            "ALTER TABLE users ALTER COLUMN is_active SET DEFAULT TRUE",
            "UPDATE users SET account_type = 'admin' WHERE role = 'admin'",
            "UPDATE users SET account_type = 'invited' WHERE account_type IS NULL",
            "UPDATE users SET is_active = TRUE WHERE is_active IS NULL",
        ]:
            try:
                conn.execute(text(statement))
                conn.commit()
            except Exception:
                conn.rollback()

        # Older invite-code tables may exist without newer admin fields.
        for statement in [
            "ALTER TABLE invite_codes ADD COLUMN token_amount INTEGER DEFAULT 1",
            "ALTER TABLE invite_codes ADD COLUMN usage_limit INTEGER DEFAULT 1",
            "ALTER TABLE invite_codes ADD COLUMN current_usage INTEGER DEFAULT 0",
            "ALTER TABLE invite_codes ADD COLUMN is_active BOOLEAN DEFAULT TRUE",
            "ALTER TABLE invite_codes ADD COLUMN created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP",
            "ALTER TABLE invite_codes ALTER COLUMN current_usage SET DEFAULT 0",
            "UPDATE invite_codes SET current_usage = 0 WHERE current_usage IS NULL",
        ]:
            try:
                conn.execute(text(statement))
                conn.commit()
            except Exception:
                conn.rollback()

    # Create default project and associate data
    with SessionLocal() as db:
        # Create a default admin user if none exists so existing data isn't completely orphaned
        default_admin = db.query(User).filter(User.email == "admin@example.com").first()
        if not default_admin:
            from passlib.context import CryptContext
            pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")
            hashed_pw = pwd_context.hash("admin")
            default_admin = User(
                email="admin@example.com",
                hashed_password=hashed_pw,
                role="admin",
                account_type="admin",
                is_active=True,
            )
            db.add(default_admin)
            db.flush()

        default_project = db.query(Project).filter(Project.name == "Default Project").first()
        if not default_project:
            default_project = Project(name="Default Project", owner_id=default_admin.id)
            db.add(default_project)
            db.flush()
        elif default_project.owner_id is None:
            default_project.owner_id = default_admin.id
            db.flush()
        
        # Ensure all existing projects are assigned to the default admin to avoid orphans
        db.query(Project).filter(Project.owner_id == None).update({Project.owner_id: default_admin.id})

        # Assign orphaned data
        db.query(Luminaire).filter(Luminaire.project_id == None).update({Luminaire.project_id: default_project.id})
        db.query(CustomColumn).filter(CustomColumn.project_id == None).update({CustomColumn.project_id: default_project.id})
        db.query(DesignTemplate).filter(DesignTemplate.project_id == None).update({DesignTemplate.project_id: default_project.id})
        
        # For AppSettings, we only migrate those that are table-specific
        project_keys = [
            "luminaire_column_order", "luminaire_hidden_columns", 
            "luminaire_sort_config", "luminaire_filter_config", 
            "luminaire_column_labels", "luminaire_filter_logic", "luminaire_row_height"
        ]
        db.query(AppSetting).filter(AppSetting.key.in_(project_keys), AppSetting.project_id == None).update({AppSetting.project_id: default_project.id})
        
        db.flush()
        for project in db.query(Project).order_by(Project.id.asc()).all():
            default_template = ensure_project_default_template(db, project.id)
            ensure_project_specification_columns(db, project.id)
            for luminaire in db.query(Luminaire).filter(Luminaire.project_id == project.id).all():
                if not str(luminaire.pdf_template or "").isdigit():
                    luminaire.pdf_template = str(default_template.id)

        db.commit()

    yield


app = FastAPI(title="Luminaire Manager", lifespan=lifespan)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:3000", "http://127.0.0.1:3000"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)
app.mount("/files", StaticFiles(directory=str(DATA_DIR)), name="files")


@contextmanager
def session_scope():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


# --- Authentication System ---
SECRET_KEY = os.getenv("SECRET_KEY", "super-secret-key-replace-me-in-production")
ALGORITHM = "HS256"
ACCESS_TOKEN_EXPIRE_MINUTES = 60 * 24 * 7 # 7 days
GOOGLE_CLIENT_ID = os.getenv("GOOGLE_CLIENT_ID", "your-google-client-id.apps.googleusercontent.com")

pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")
oauth2_scheme = OAuth2PasswordBearer(tokenUrl="/api/auth/login")

def verify_password(plain_password, hashed_password):
    if not hashed_password: return False
    return pwd_context.verify(plain_password, hashed_password)

def get_password_hash(password):
    return pwd_context.hash(password)

def create_access_token(data: dict, expires_delta: Optional[timedelta] = None):
    to_encode = data.copy()
    if expires_delta:
        expire = datetime.utcnow() + expires_delta
    else:
        expire = datetime.utcnow() + timedelta(minutes=15)
    to_encode.update({"exp": expire})
    encoded_jwt = jwt.encode(to_encode, SECRET_KEY, algorithm=ALGORITHM)
    return encoded_jwt

async def get_current_user(token: str = Depends(oauth2_scheme), db: Session = Depends(get_db)):
    credentials_exception = HTTPException(
        status_code=401,
        detail="Could not validate credentials",
        headers={"WWW-Authenticate": "Bearer"},
    )
    try:
        payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
        user_id: str = str(payload.get("sub"))
        if user_id is None:
            raise credentials_exception
    except JWTError:
        raise credentials_exception
    user = db.query(User).filter(User.id == int(user_id)).first()
    if user is None:
        raise credentials_exception
    if not user.is_active:
        raise HTTPException(status_code=403, detail="Your account is pending admin activation")
    return user

async def get_admin_user(current_user: User = Depends(get_current_user)):
    if current_user.role != "admin":
        raise HTTPException(status_code=403, detail="Not enough privileges")
    return current_user

class SignupRequest(BaseModel):
    email: str
    password: str
    invite_code: Optional[str] = None
    name: Optional[str] = None
    studio_name: Optional[str] = None

class GoogleAuthRequest(BaseModel):
    credential: str


def serialize_user_session(row: User) -> dict:
    return {
        "id": row.id,
        "email": row.email,
        "role": row.role,
        "tokens": row.tokens or 0,
        "accountType": row.account_type,
        "isActive": bool(row.is_active),
        "name": row.name or None,
        "studioName": row.studio_name or None,
    }

@app.post("/api/auth/signup")
def signup(request: SignupRequest, db: Session = Depends(get_db)):
    email = request.email.strip().lower()
    user = db.query(User).filter(User.email == email).first()
    if user:
        raise HTTPException(status_code=400, detail="Email already registered")
    
    # First user becomes admin. Later users can optionally join through an invite code.
    is_first = db.query(User).count() == 0
    role = "admin" if is_first else "client"
    invite_code = None
    account_type = "admin" if is_first else "freemium"
    is_active = is_first
    starting_tokens = 0 if is_first else FREEMIUM_DEFAULT_CREDITS
    if request.invite_code and not is_first:
        code_value = request.invite_code.strip().upper()
        invite_code = db.query(InviteCode).filter(InviteCode.code == code_value).first()
        if not invite_code or not invite_code.is_active:
            raise HTTPException(status_code=400, detail="Invalid invite code")
        current_usage = db.query(User).filter(User.invite_code_id == invite_code.id).count()
        if current_usage >= invite_code.usage_limit:
            raise HTTPException(status_code=400, detail="Invite code has reached its user limit")
        starting_tokens = invite_code.token_amount
        account_type = "invited"
        is_active = True
    
    hashed_pass = get_password_hash(request.password)
    new_user = User(
        email=email,
        hashed_password=hashed_pass,
        role=role,
        account_type=account_type,
        is_active=is_active,
        tokens=starting_tokens,
        invite_code_id=invite_code.id if invite_code else None,
        name=(request.name or "").strip()[:120] or None,
        studio_name=(request.studio_name or "").strip()[:120] or None,
    )
    db.add(new_user)
    db.commit()
    db.refresh(new_user)

    if not new_user.is_active:
        return {
            "requires_activation": True,
            "message": f"Freemium account created with {FREEMIUM_DEFAULT_CREDITS} PDF credits. An administrator must activate it before you can sign in.",
            "user": serialize_user_session(new_user),
        }
    
    access_token = create_access_token(
        data={"sub": str(new_user.id), "role": new_user.role},
        expires_delta=timedelta(minutes=ACCESS_TOKEN_EXPIRE_MINUTES)
    )
    return {"access_token": access_token, "token_type": "bearer", "user": serialize_user_session(new_user)}

@app.post("/api/auth/login")
def login(request: OAuth2PasswordRequestForm = Depends(), db: Session = Depends(get_db)):
    user = db.query(User).filter(User.email == request.username.strip().lower()).first()
    if not user or not verify_password(request.password, user.hashed_password):
        raise HTTPException(status_code=401, detail="Incorrect email or password")
    if not user.is_active:
        raise HTTPException(status_code=403, detail="Your freemium account is awaiting admin activation")
    
    access_token = create_access_token(
        data={"sub": str(user.id), "role": user.role},
        expires_delta=timedelta(minutes=ACCESS_TOKEN_EXPIRE_MINUTES)
    )
    return {"access_token": access_token, "token_type": "bearer", "user": serialize_user_session(user)}

@app.post("/api/auth/google")
def google_auth(request: GoogleAuthRequest, db: Session = Depends(get_db)):
    try:
        idinfo = id_token.verify_oauth2_token(request.credential, google_requests.Request(), GOOGLE_CLIENT_ID)
        email = idinfo['email'].strip().lower()
        google_sub = idinfo['sub']
        
        user = db.query(User).filter(User.email == email).first()
        if not user:
            is_first = db.query(User).count() == 0
            role = "admin" if is_first else "client"
            user = User(
                email=email.strip().lower(),
                google_sub=google_sub,
                role=role,
                account_type="admin" if is_first else "freemium",
                is_active=is_first,
                tokens=0 if is_first else FREEMIUM_DEFAULT_CREDITS,
            )
            db.add(user)
            db.commit()
            db.refresh(user)
        elif not user.google_sub:
            # Link existing account to google
            user.google_sub = google_sub
            db.commit()

        if not user.is_active:
            return {
                "requires_activation": True,
                "message": f"Freemium account created with {FREEMIUM_DEFAULT_CREDITS} PDF credits. An administrator must activate it before you can sign in.",
                "user": serialize_user_session(user),
            }
            
        access_token = create_access_token(
            data={"sub": str(user.id), "role": user.role},
            expires_delta=timedelta(minutes=ACCESS_TOKEN_EXPIRE_MINUTES)
        )
        return {"access_token": access_token, "token_type": "bearer", "user": serialize_user_session(user)}
    except ValueError:
        raise HTTPException(status_code=401, detail="Invalid Google Token")

def invite_code_usage(db: Session, invite_code_id: int) -> int:
    return db.query(User).filter(User.invite_code_id == invite_code_id).count()


def serialize_invite_code(db: Session, row: InviteCode) -> dict:
    current_usage = invite_code_usage(db, row.id)
    return {
        "id": row.id,
        "code": row.code,
        "tokenAmount": row.token_amount,
        "usageLimit": row.usage_limit,
        "currentUsage": current_usage,
        "isActive": bool(row.is_active),
        "spotsLeft": max(row.usage_limit - current_usage, 0),
    }


def serialize_user_admin(row: User) -> dict:
    invite = row.invite_code
    return {
        "id": row.id,
        "email": row.email,
        "role": row.role,
        "accountType": row.account_type,
        "isActive": bool(row.is_active),
        "tokens": row.tokens or 0,
        "inviteCode": invite.code if invite else None,
        "inviteCodeId": invite.id if invite else None,
        "inviteTokenAmount": invite.token_amount if invite else None,
        "createdAt": row.created_at.isoformat(),
        "name": row.name or None,
        "studioName": row.studio_name or None,
    }


def reserve_freemium_credits(db: Session, user_id: int, requested_credits: int) -> Optional[int]:
    """Atomically reserve credits before extraction and return the remaining balance."""
    if requested_credits <= 0:
        return None

    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=401, detail="Could not validate credentials")
    if not user.is_active:
        raise HTTPException(status_code=403, detail="Your account is pending admin activation")
    if user.role == "admin" or user.account_type != "freemium":
        return None

    updated = (
        db.query(User)
        .filter(
            User.id == user_id,
            User.is_active.is_(True),
            User.account_type == "freemium",
            User.tokens >= requested_credits,
        )
        .update({User.tokens: User.tokens - requested_credits}, synchronize_session=False)
    )
    if updated != 1:
        db.rollback()
        available = db.query(User.tokens).filter(User.id == user_id).scalar() or 0
        raise HTTPException(
            status_code=402,
            detail=f"Not enough PDF credits. This request needs {requested_credits}, but you have {available} remaining.",
        )

    db.commit()
    return db.query(User.tokens).filter(User.id == user_id).scalar() or 0


@app.get("/api/me")
def get_me(current_user: User = Depends(get_current_user)):
    return serialize_user_session(current_user)


@app.get("/api/users")
def get_users(db: Session = Depends(get_db), current_user: User = Depends(get_admin_user)):
    users = db.query(User).order_by(User.id.asc()).all()
    return {"users": [serialize_user_admin(u) for u in users]}


@app.patch("/api/admin/users/{user_id}")
async def update_admin_user(
    user_id: int,
    request: Request,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_admin_user),
):
    row = db.query(User).filter(User.id == user_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="User not found")
    if row.role == "admin":
        raise HTTPException(status_code=400, detail="Administrator accounts cannot be changed here")

    data = await request.json()

    if "is_active" in data:
        if not isinstance(data["is_active"], bool):
            raise HTTPException(status_code=400, detail="is_active must be true or false")
        row.is_active = data["is_active"]

    if "set_tokens" in data:
        try:
            set_tokens = int(data["set_tokens"])
        except (TypeError, ValueError):
            raise HTTPException(status_code=400, detail="Credits must be a whole number")
        if set_tokens < 0:
            raise HTTPException(status_code=400, detail="Credits cannot be negative")
        row.tokens = set_tokens

    if "top_up_tokens" in data:
        try:
            top_up_tokens = int(data["top_up_tokens"])
        except (TypeError, ValueError):
            raise HTTPException(status_code=400, detail="Credit top-up must be a whole number")
        if top_up_tokens < 1:
            raise HTTPException(status_code=400, detail="Credit top-up must be at least 1")
        row.tokens = (row.tokens or 0) + top_up_tokens

    if "invite_code_id" in data:
        invite_code_id = data["invite_code_id"]
        if invite_code_id is None:
            row.invite_code_id = None
        else:
            try:
                invite_code_id = int(invite_code_id)
            except (TypeError, ValueError):
                raise HTTPException(status_code=400, detail="Invalid invite code")
            invite = db.query(InviteCode).filter(InviteCode.id == invite_code_id).first()
            if not invite:
                raise HTTPException(status_code=404, detail="Invite code not found")
            if row.invite_code_id != invite.id and invite_code_usage(db, invite.id) >= invite.usage_limit:
                raise HTTPException(status_code=400, detail="Invite code has reached its user limit")
            row.invite_code_id = invite.id

    db.commit()
    db.refresh(row)
    return serialize_user_admin(row)


@app.get("/api/admin/invite-codes")
def list_invite_codes(db: Session = Depends(get_db), current_user: User = Depends(get_admin_user)):
    rows = db.query(InviteCode).order_by(InviteCode.id.asc()).all()
    return {"codes": [serialize_invite_code(db, row) for row in rows]}


@app.post("/api/admin/invite-codes")
async def create_invite_code(request: Request, db: Session = Depends(get_db), current_user: User = Depends(get_admin_user)):
    data = await request.json()
    code = str(data.get("code") or "").strip().upper()
    token_amount = int(data.get("token_amount") or data.get("tokenAmount") or 0)
    usage_limit = int(data.get("usage_limit") or data.get("usageLimit") or 0)
    if not code:
        raise HTTPException(status_code=400, detail="Invite code is required")
    if token_amount < 1:
        raise HTTPException(status_code=400, detail="Credits must be at least 1")
    if usage_limit < 1:
        raise HTTPException(status_code=400, detail="User limit must be at least 1")
    if db.query(InviteCode).filter(InviteCode.code == code).first():
        raise HTTPException(status_code=400, detail="Invite code already exists")
    row = InviteCode(code=code, token_amount=token_amount, usage_limit=usage_limit, is_active=True)
    db.add(row)
    db.commit()
    db.refresh(row)
    return {"code": serialize_invite_code(db, row)}


@app.patch("/api/admin/invite-codes/{code_id}")
async def update_invite_code(code_id: int, request: Request, db: Session = Depends(get_db), current_user: User = Depends(get_admin_user)):
    row = db.query(InviteCode).filter(InviteCode.id == code_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Invite code not found")
    data = await request.json()
    current_usage = invite_code_usage(db, row.id)
    if "token_amount" in data or "tokenAmount" in data:
        token_amount = int(data.get("token_amount") or data.get("tokenAmount") or 0)
        if token_amount < 1:
            raise HTTPException(status_code=400, detail="Credits must be at least 1")
        row.token_amount = token_amount
    if "usage_limit" in data or "usageLimit" in data:
        usage_limit = int(data.get("usage_limit") or data.get("usageLimit") or 0)
        if usage_limit < current_usage:
            raise HTTPException(status_code=400, detail=f"User limit cannot be less than current usage ({current_usage})")
        row.usage_limit = usage_limit
    if "is_active" in data:
        row.is_active = bool(data.get("is_active"))
    db.commit()
    db.refresh(row)
    return {"code": serialize_invite_code(db, row)}


@app.delete("/api/admin/invite-codes/{code_id}")
def delete_invite_code(code_id: int, db: Session = Depends(get_db), current_user: User = Depends(get_admin_user)):
    row = db.query(InviteCode).filter(InviteCode.id == code_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Invite code not found")
    db.query(User).filter(User.invite_code_id == row.id).update({User.invite_code_id: None})
    db.delete(row)
    db.commit()
    return {"status": "ok"}


def get_project_or_404(db: Session, project_id: int, current_user: User) -> Project:
    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    if current_user.role != "admin" and project.owner_id != current_user.id:
        raise HTTPException(status_code=403, detail="Not authorized to access this project")
    return project

def to_public_file_url(relative_path: Optional[str]) -> Optional[str]:
    if not relative_path:
        return None
    return f"/files/{relative_path.replace(os.sep, '/')}"


def get_setting(db: Session, key: str, default: str = "", project_id: Optional[int] = None) -> str:
    query = db.query(AppSetting).filter(AppSetting.key == key)
    if project_id is not None:
        query = query.filter(AppSetting.project_id == project_id)
    else:
        query = query.filter(AppSetting.project_id == None)
    row = query.first()
    return row.value if row and row.value else default


def set_setting(db: Session, key: str, value: str, project_id: Optional[int] = None) -> None:
    query = db.query(AppSetting).filter(AppSetting.key == key)
    if project_id is not None:
        query = query.filter(AppSetting.project_id == project_id)
    else:
        query = query.filter(AppSetting.project_id == None)
    
    row = query.first()
    if not row:
        row = AppSetting(key=key, value=value, project_id=project_id)
        db.add(row)
    else:
        row.value = value
        row.updated_at = datetime.utcnow()


def is_default_specification_template(template: DesignTemplate) -> bool:
    design = template.design_json or {}
    if isinstance(design, str):
        try:
            design = json.loads(design)
        except Exception:
            return False
    return design.get("systemTemplateKey") == DEFAULT_TEMPLATE_KEY


def ensure_project_default_template(db: Session, project_id: int) -> DesignTemplate:
    current_design = build_default_specification_template_design()
    templates = (
        db.query(DesignTemplate)
        .filter(DesignTemplate.project_id == project_id)
        .order_by(DesignTemplate.id.asc())
        .all()
    )
    for template in templates:
        if is_default_specification_template(template):
            design = template.design_json or {}
            if isinstance(design, str):
                try:
                    design = json.loads(design)
                except Exception:
                    design = {}
            try:
                installed_version = int(design.get("systemTemplateVersion", 0))
            except (TypeError, ValueError):
                installed_version = 0
            if installed_version < DEFAULT_TEMPLATE_VERSION:
                template.page_size = "A4"
                template.page_width = DEFAULT_TEMPLATE_WIDTH
                template.page_height = DEFAULT_TEMPLATE_HEIGHT
                template.design_json = current_design
                template.updated_at = datetime.utcnow()
            return template

    template = DesignTemplate(
        project_id=project_id,
        luminaire_id=None,
        name=DEFAULT_TEMPLATE_NAME,
        page_size="A4",
        page_width=DEFAULT_TEMPLATE_WIDTH,
        page_height=DEFAULT_TEMPLATE_HEIGHT,
        design_json=current_design,
    )
    db.add(template)
    db.flush()
    return template


# --- RISE OS studio + project brand -------------------------------------------
BRAND_TEMPLATES = ("standard", "compact")
DEFAULT_BRAND_ACCENT = DEFAULT_TEMPLATE_ACCENT.lower()
_HEX_COLOR_RE = re.compile(r"^#[0-9a-fA-F]{6}$")
STUDIO_LOGO_MAX_BYTES = 2 * 1024 * 1024


def normalize_accent(value: Optional[str]) -> str:
    value = str(value or "").strip()
    return value.lower() if _HEX_COLOR_RE.match(value) else ""


def normalize_brand_template(value: Optional[str]) -> str:
    value = str(value or "").strip().lower()
    return value if value in BRAND_TEMPLATES else "standard"


def studio_setting_key(user_id: int, name: str) -> str:
    return f"studio:{user_id}:{name}"


def get_studio_defaults(db: Session, user: "User") -> dict:
    logo_path = get_setting(db, studio_setting_key(user.id, "logo_path"), "")
    if logo_path and not (DATA_DIR / logo_path).exists():
        logo_path = ""
    name = (user.studio_name or "").strip()
    if not name and user.role == "admin":
        name = get_setting(db, "company_name", "")
    return {
        "studioName": name,
        "accent": normalize_accent(get_setting(db, studio_setting_key(user.id, "accent"), "")) or DEFAULT_BRAND_ACCENT,
        "template": normalize_brand_template(get_setting(db, studio_setting_key(user.id, "template"), "standard")),
        "logoPath": logo_path,
        "logoUrl": to_public_file_url(logo_path) or "",
        "saved": get_setting(db, studio_setting_key(user.id, "saved"), "") == "1",
    }


def get_project_brand(db: Session, project: "Project") -> dict:
    pid = project.id
    return {
        "projectId": pid,
        "name": project.name,
        "location": get_setting(db, "specification_location_country", "", project_id=pid),
        "revision": get_setting(db, "specification_revision", "A", project_id=pid),
        "issueDate": get_setting(db, "specification_issue_date", "", project_id=pid),
        "studioName": get_setting(db, "brand_studio_name", "", project_id=pid),
        "accent": normalize_accent(get_setting(db, "brand_accent", "", project_id=pid)) or DEFAULT_BRAND_ACCENT,
        "template": normalize_brand_template(get_setting(db, "brand_template", "standard", project_id=pid)),
        "logoUrl": project_logo_url(db, pid),
        "hasOwnLogo": bool(get_setting(db, "specification_logo_path", "", project_id=pid)),
    }


def is_compact_specification_template(template: DesignTemplate) -> bool:
    design = template.design_json or {}
    if isinstance(design, str):
        try:
            design = json.loads(design)
        except Exception:
            return False
    return design.get("systemTemplateKey") == COMPACT_TEMPLATE_KEY


def ensure_project_compact_template(db: Session, project_id: int) -> DesignTemplate:
    for template in (
        db.query(DesignTemplate).filter(DesignTemplate.project_id == project_id).order_by(DesignTemplate.id.asc()).all()
    ):
        if is_compact_specification_template(template):
            return template
    template = DesignTemplate(
        project_id=project_id,
        luminaire_id=None,
        name=COMPACT_TEMPLATE_NAME,
        page_size="A4",
        page_width=DEFAULT_TEMPLATE_WIDTH,
        page_height=DEFAULT_TEMPLATE_HEIGHT,
        design_json=build_compact_specification_template_design(),
    )
    db.add(template)
    db.flush()
    return template


def project_brand_template(db: Session, project_id: int) -> DesignTemplate:
    """Template new luminaires in this project should use by default."""
    choice = normalize_brand_template(get_setting(db, "brand_template", "standard", project_id=project_id))
    if choice == "compact":
        return ensure_project_compact_template(db, project_id)
    return ensure_project_default_template(db, project_id)


def switch_project_brand_template(db: Session, project_id: int, new_choice: str) -> None:
    """Point luminaires that used the previous project template at the new one.
    Generated documents are snapshots and are not touched."""
    old_template = project_brand_template(db, project_id)
    set_setting(db, "brand_template", normalize_brand_template(new_choice), project_id=project_id)
    db.flush()
    new_template = project_brand_template(db, project_id)
    if old_template.id != new_template.id:
        db.query(Luminaire).filter(
            Luminaire.project_id == project_id,
            Luminaire.pdf_template == str(old_template.id),
        ).update({Luminaire.pdf_template: str(new_template.id)}, synchronize_session=False)


def branded_item_style(style: dict, accent: str) -> dict:
    """Swap the template's system accent colour for the project's accent."""
    if not accent or accent == DEFAULT_BRAND_ACCENT:
        return style
    out = dict(style)
    for key in ("backgroundColor", "borderColor", "color"):
        value = str(out.get(key) or "").strip().lower()
        if value == DEFAULT_BRAND_ACCENT:
            out[key] = accent
    return out


def save_brand_logo(content: bytes, prefix: str) -> str:
    if not content:
        raise HTTPException(status_code=400, detail="Logo file is empty")
    if len(content) > STUDIO_LOGO_MAX_BYTES:
        raise HTTPException(status_code=400, detail="Choose a PNG or JPEG smaller than 2 MB.")
    try:
        from PIL import Image

        with Image.open(BytesIO(content)) as source:
            source.verify()
            fmt = (source.format or "").upper()
        if fmt not in {"PNG", "JPEG"}:
            raise HTTPException(status_code=400, detail="Choose a PNG or JPEG smaller than 2 MB.")
        with Image.open(BytesIO(content)) as source:
            if source.width > 6000 or source.height > 6000:
                raise HTTPException(status_code=400, detail="Logo dimensions are too large")
            relative_path = f"branding/{prefix}-{uuid.uuid4().hex}.png"
            source.convert("RGBA").save(DATA_DIR / relative_path, "PNG", optimize=True)
            return relative_path
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=400, detail="Unsupported or invalid image file")


def ensure_project_specification_columns(db: Session, project_id: int) -> list[CustomColumn]:
    existing_columns = (
        db.query(CustomColumn)
        .filter(CustomColumn.project_id == project_id)
        .order_by(CustomColumn.id.asc())
        .all()
    )
    by_key = {
        canonicalize_import_key(column.name): column
        for column in existing_columns
    }
    saved_order = get_saved_column_order(db, project_id)
    changed_order = False
    if not saved_order and not existing_columns:
        # Brand-new project: core columns (Type, Photo, ...) first, specification
        # columns after them, matching the established project layout.
        saved_order = list(SCHEDULE_EXPORT_LABELS.keys())
        changed_order = True
    ensured: list[CustomColumn] = []

    for specification in DEFAULT_SPECIFICATION_COLUMNS:
        accepted_keys = {
            canonicalize_import_key(specification["name"]),
            *(
                canonicalize_import_key(alias)
                for alias in specification.get("aliases", ())
            ),
        }
        column = next(
            (by_key[key] for key in accepted_keys if key in by_key),
            None,
        )
        if column is None:
            column = CustomColumn(
                project_id=project_id,
                name=specification["name"],
                ai_enabled=True,
                ai_description=specification["description"],
            )
            db.add(column)
            db.flush()
            existing_columns.append(column)
            by_key[canonicalize_import_key(column.name)] = column
        else:
            column.ai_enabled = True
            column.ai_description = specification["description"]

        column_key = f"custom:{column.id}"
        if column_key not in saved_order:
            saved_order.append(column_key)
            changed_order = True
        ensured.append(column)

    if changed_order:
        save_column_order_setting(db, saved_order, project_id)
    return ensured


def project_logo_url(db: Session, project_id: int) -> str:
    configured_path = get_setting(db, "specification_logo_path", "", project_id=project_id)
    if configured_path and (DATA_DIR / configured_path).exists():
        return to_public_file_url(configured_path) or ""
    if get_setting(db, "brand_hide_default_logo", "", project_id=project_id) == "1":
        return ""
    if SCHEDULE_BRAND_LOGO_SOURCE_PATH.exists():
        return to_public_file_url("branding/afterglow-logo.jpeg") or ""
    return ""


# --- PDF Export Utilities ---
def hex_to_reportlab(color: Optional[str], fallback=colors.transparent):
    if not color: return fallback
    val = color.strip().lstrip("#")
    if len(val) == 3: val = "".join(c*2 for c in val)
    if len(val) != 6: return fallback
    try: return colors.HexColor(f"#{val}")
    except: return fallback

def get_internal_file_path(url: Optional[str]) -> Optional[Path]:
    if not url or not url.startswith("/files/"): return None
    return DATA_DIR / url.removeprefix("/files/")

def resolve_asset_url(db: Session, row: Luminaire, attr_name: str) -> str:
    path = getattr(row, attr_name)
    if path and (DATA_DIR / path).exists(): return to_public_file_url(path)
    return ""

def designer_context_from_luminaire(db: Session, row: Luminaire) -> dict:
    payload = serialize_luminaire(row)
    custom = payload.get("customData") or {}
    normalized_custom = {
        canonicalize_import_key(key): clean_optional_text(value)
        for key, value in custom.items()
    }

    def custom_value(*aliases: str) -> str:
        for alias in aliases:
            value = normalized_custom.get(canonicalize_import_key(alias))
            if value:
                return value
        return ""

    project = db.query(Project).filter(Project.id == row.project_id).first()
    project_id = int(row.project_id or 0)
    luminaire_reference = (
        custom_value("Luminaire Ref", "Luminaire Reference")
        or clean_optional_text(row.spec_code)
        or extract_luminaire_reference_from_source_name(row.source_name)
        or ""
    )
    context = {
        **payload,
        **custom,
        "projectName": project.name if project else "",
        "projectLogoUrl": project_logo_url(db, project_id) if project_id else "",
        "projectLocationCountry": get_setting(db, "specification_location_country", "", project_id=project_id),
        "projectRevision": get_setting(db, "specification_revision", "A (Current)", project_id=project_id),
        "projectIssueDate": get_setting(db, "specification_issue_date", "", project_id=project_id) or datetime.utcnow().date().isoformat(),
        "studioName": get_setting(db, "brand_studio_name", "", project_id=project_id) if project_id else "",
        "luminaireReference": luminaire_reference,
        "luminaireLocation": custom_value("Location", "Luminaire Location"),
        "productCode": clean_optional_text(row.spec_code) or luminaire_reference,
        "productName": clean_optional_text(row.product_name) or "",
        "mountingType": custom_value("Mounting Type", "Mounting") or clean_optional_text(row.type_code) or "",
        "dimming": clean_optional_text(row.dimming) or custom_value("Dimming / Control", "Control"),
        "lightSource": custom_value("Light Source", "Lamp Type", "LED Type"),
        "wattage": clean_optional_text(row.unit_watts) or custom_value("Wattage", "Power"),
        "colourTemperature": clean_optional_text(row.cct) or custom_value("Colour Temperature", "Color Temperature"),
        "macadam": custom_value("MacAdam / SDCM", "MacAdam", "SDCM"),
        "beamDistribution": custom_value("Beam / Distribution", "Beam", "Beam Angle", "Distribution"),
        "deliveredLumens": custom_value("Delivered Lumens", "Luminous Flux", "Lumens"),
        "ledLifetime": custom_value("LED Lifetime", "Lifetime"),
        "inputVoltage": custom_value("Input Voltage", "Voltage"),
        "emergencyProvision": custom_value("Emergency Provision", "Emergency"),
        "dimensions": custom_value("Dimensions (LxWxH / ØxH)", "Dimensions", "Dimension"),
        "material": custom_value("Material", "Housing Material"),
        "finishColour": clean_optional_text(row.colour) or custom_value("Finish / Colour", "Finish", "Colour"),
        "ikRating": custom_value("IK Rating", "IK"),
        "weight": custom_value("Weight"),
        "operatingTemperature": custom_value("Operating Temp.", "Operating Temperature"),
        "requiredAccessories": custom_value("Required Accessories", "Accessories"),
        "generalNotes": custom_value("General Notes", "Notes"),
        "sourceDocument": clean_optional_text(row.source_name) or "",
        "productImageUrl": resolve_asset_url(db, row, "product_image_path"),
        "diagramImageUrl": resolve_asset_url(db, row, "diagram_image_path"),
        "polarImageUrl": resolve_asset_url(db, row, "polar_image_path"),
    }
    return {
        **context,
    }

def page_dimensions(page_size: str, width: int, height: int) -> tuple[float, float]:
    if page_size == "LETTER": return LETTER
    if page_size == "CUSTOM": return float(width), float(height)
    return A4

def render_text_item(pdf: canvas.Canvas, item: dict, value: str, page_height: float):
    x, y = float(item.get("x", 0)), page_height - float(item.get("y", 0)) - float(item.get("height", 0))
    w, h = float(item.get("width", 180)), float(item.get("height", 50))
    style = item.get("style", {})
    font = "Helvetica-Bold" if style.get("fontWeight") == "bold" else "Helvetica"
    size = float(style.get("fontSize", 12))
    pdf.setFillColor(hex_to_reportlab(style.get("backgroundColor"), colors.transparent))
    pdf.rect(x, y, w, h, stroke=0, fill=1)
    bc = hex_to_reportlab(style.get("borderColor"), colors.transparent)
    bw = float(style.get("borderWidth", 0))
    if bc != colors.transparent and bw > 0:
        pdf.setStrokeColor(bc)
        pdf.setLineWidth(bw)
        pdf.rect(x, y, w, h, stroke=1, fill=0)
    pdf.setFillColor(hex_to_reportlab(style.get("color"), colors.black))
    pdf.setFont(font, size)
    pad = float(style.get("padding", 6))
    lines = simpleSplit(str(value or ""), font, size, max(20, w - (pad * 2)))
    line_height = size + 2
    lines = lines[: max(1, int(max(1, h - (pad * 2)) // line_height))]
    vertical_align = style.get("verticalAlign", "top")
    if vertical_align == "middle":
        block_height = max(size, (len(lines) * line_height) - 2)
        baseline = y + ((h + block_height) / 2) - size
    elif vertical_align == "bottom":
        baseline = y + pad + ((len(lines) - 1) * line_height)
    else:
        baseline = y + h - pad - size
    text_align = style.get("textAlign", "left")
    for line in lines:
        if text_align == "center":
            pdf.drawCentredString(x + (w / 2), baseline, line)
        elif text_align == "right":
            pdf.drawRightString(x + w - pad, baseline, line)
        else:
            pdf.drawString(x + pad, baseline, line)
        baseline -= line_height

def render_image_item(pdf: canvas.Canvas, item: dict, img_path: Optional[Path], page_height: float):
    if not img_path or not img_path.exists(): return
    x, y = float(item.get("x", 0)), page_height - float(item.get("y", 0)) - float(item.get("height", 120))
    w, h = float(item.get("width", 120)), float(item.get("height", 120))
    pad = float((item.get("style") or {}).get("padding", 0))
    x, y, w, h = x + pad, y + pad, max(1, w - (pad * 2)), max(1, h - (pad * 2))
    try:
        reader = ImageReader(str(img_path))
        image_width, image_height = reader.getSize()
        scale = min(w / image_width, h / image_height)
        draw_width, draw_height = image_width * scale, image_height * scale
        pdf.drawImage(
            reader,
            x + ((w - draw_width) / 2),
            y + ((h - draw_height) / 2),
            width=draw_width,
            height=draw_height,
            preserveAspectRatio=True,
            mask="auto",
        )
    except Exception:
        return

def create_pdf_export(db: Session, row: Luminaire, template: DesignTemplate) -> str:
    pw, ph = page_dimensions(template.page_size, template.page_width, template.page_height)
    name = f"{uuid.uuid4().hex}.pdf"
    rel = f"exports/{name}"
    pdf = canvas.Canvas(str(EXPORTS_DIR / name), pagesize=(pw, ph))
    data = designer_context_from_luminaire(db, row)
    items = (template.design_json or {}).get("items", [])
    accent = normalize_accent(get_setting(db, "brand_accent", "", project_id=row.project_id)) if row.project_id else ""
    if accent and accent != DEFAULT_BRAND_ACCENT:
        items = [{**item, "style": branded_item_style(item.get("style") or {}, accent)} for item in items]
    for item in items:
        it = item.get("type")
        if it == "image" and item.get("bindingKey") == "projectLogoUrl" and not data.get("projectLogoUrl") and data.get("studioName"):
            render_text_item(pdf, {**item, "style": {**(item.get("style") or {}), "fontSize": 9, "fontWeight": "bold", "textAlign": "center", "verticalAlign": "middle"}}, data["studioName"], ph)
            continue
        if it in {"field", "text"}:
            val = data.get(item.get("bindingKey", ""), "") if it == "field" else item.get("text", "")
            if it == "field" and (val is None or str(val).strip() == ""):
                val = item.get("emptyText", "Not available")
            render_text_item(pdf, item, val, ph)
        elif it == "image":
            ip = get_internal_file_path(item.get("imageUrl") or data.get(item.get("bindingKey", ""), ""))
            render_image_item(pdf, item, ip, ph)
        elif it in {"shape", "line"}:
            x, y = float(item.get("x", 0)), ph - float(item.get("y", 0)) - float(item.get("height", 20))
            w, h = float(item.get("width", 100)), float(item.get("height", 20))
            style = item.get("style", {})
            pdf.setFillColor(hex_to_reportlab(style.get("backgroundColor"), colors.transparent))
            pdf.rect(x, y, w, h, stroke=0, fill=1)
            bc = hex_to_reportlab(style.get("borderColor"), colors.transparent)
            bw = float(style.get("borderWidth", 0))
            if bc != colors.transparent and bw > 0:
                pdf.setStrokeColor(bc)
                pdf.setLineWidth(bw)
                pdf.rect(x, y, w, h, stroke=1, fill=0)
    pdf.showPage()
    pdf.save()
    return rel

def serialize_project(row: Project) -> dict:
    return {
        "id": row.id,
        "name": row.name,
        "createdAt": row.created_at.isoformat(),
    }

def serialize_template(row: DesignTemplate) -> dict:
    design = row.design_json
    if isinstance(design, str):
        try: design = json.loads(design)
        except: design = {"items": []}
    return {
        "id": row.id,
        "sampleLuminaireId": row.luminaire_id,
        "name": row.name,
        "pageSize": row.page_size,
        "pageWidth": row.page_width,
        "pageHeight": row.page_height,
        "designJson": design,
        "isDefault": design.get("systemTemplateKey") == DEFAULT_TEMPLATE_KEY,
        "updatedAt": row.updated_at.isoformat()
    }

def serialize_generated_document(row: GeneratedDocument) -> dict:
    return {
        "id": row.id,
        "fileUrl": to_public_file_url(row.file_path),
        "title": row.title,
        "createdAt": row.created_at.isoformat()
    }



def serialize_luminaire(row: Luminaire) -> dict:
    custom = {}
    if row.custom_data:
        try:
            custom = json.loads(row.custom_data)
        except Exception:
            pass

    return {
        "id": row.id,
        "sourceType": row.source_type,
        "sourceName": row.source_name,
        "sourceUrl": row.source_url,
        "documentUrl": to_public_file_url(row.document_path),
        "type": row.type_code,
        "manufacturer": normalize_manufacturer_name(row.manufacturer),
        "product": row.product_name,
        "specCode": row.spec_code,
        "description": row.description,
        "colour": row.colour,
        "cct": row.cct,
        "cri": row.cri,
        "unitWatts": row.unit_watts,
        "dimming": row.dimming,
        "quantity": row.quantity,
        "ipRating": row.ip_rating,
        "totalWatts": row.total_watts,
        "diagramImageUrl": to_public_file_url(row.diagram_image_path),
        "productImageUrl": to_public_file_url(row.product_image_path),
        "polarImageUrl": to_public_file_url(row.polar_image_path),
        "customData": custom,
        "pdfTemplate": row.pdf_template,
    }


def dedupe_ints_preserve_order(values: List[int]) -> List[int]:
    ordered: List[int] = []
    seen = set()
    for raw in values:
        try:
            value = int(raw)
        except Exception:
            continue
        if value in seen:
            continue
        seen.add(value)
        ordered.append(value)
    return ordered


def canonicalize_import_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())


def clean_optional_text(value: object) -> Optional[str]:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value).strip())
    return text or None


STANDARD_COLUMN_NAME_KEYS = {
    canonicalize_import_key(label) for label in SCHEDULE_EXPORT_LABELS.values()
} | {
    "quantity",
    "quanity",
    "qty",
    "watts",
    "unitwatt",
    "unitwatts",
    "totalwatt",
    "totalwatts",
    "template",
    "templatepdf",
}


def is_standard_column_name(name: object) -> bool:
    return canonicalize_import_key(name) in STANDARD_COLUMN_NAME_KEYS


def normalize_manufacturer_name(value: object) -> Optional[str]:
    text = clean_optional_text(value)
    if not text:
        return None

    if text.lower() != text and text.upper() != text:
        return text

    tokens = []
    for token in text.lower().split(" "):
        if re.fullmatch(r"[a-z]{1,4}", token):
            tokens.append(token.upper())
            continue
        tokens.append(
            re.sub(
                r"[A-Za-zÀ-ÿ]+",
                lambda match: match.group(0)[:1].upper() + match.group(0)[1:].lower(),
                token,
            )
        )
    return " ".join(tokens)


def normalize_type_name(
    value: object,
    *,
    product: object = None,
    description: object = None,
) -> Optional[str]:
    text = clean_optional_text(value)
    product_text = clean_optional_text(product) or ""
    description_text = clean_optional_text(description) or ""
    combined_context = f"{product_text} {description_text}".strip().lower()

    if not text:
        if "flex" in combined_context:
            return "Flex"
        return None

    lowered = text.lower()
    if lowered in {"luminaire", "led", "lighting"} and "flex" in combined_context:
        return "Flex"

    if text.lower() != text and text.upper() != text:
        return text

    return re.sub(
        r"[A-Za-zÀ-ÿ]+",
        lambda match: match.group(0)[:1].upper() + match.group(0)[1:].lower(),
        text,
    )


def parse_decimal_from_text(value: object) -> Optional[Decimal]:
    text = clean_optional_text(value)
    if not text:
        return None
    match = re.search(r"[-+]?\d+(?:[.,]\d+)?", text)
    if not match:
        return None
    try:
        return Decimal(match.group(0).replace(",", "."))
    except InvalidOperation:
        return None


def format_watts_total(value: Decimal) -> str:
    quantized = value.normalize()
    if quantized == quantized.to_integral():
        return f"{int(quantized)}W"
    return f"{format(quantized, 'f')}W"


def has_multiple_wattage_options(value: object) -> bool:
    text = clean_optional_text(value)
    if not text:
        return False
    watt_matches = re.findall(r"\d+(?:[.,]\d+)?\s*W\b", text, flags=re.I)
    if len(watt_matches) > 1:
        return True
    if re.search(r"\d+(?:[.,]\d+)?\s*(?:/|,|;|\|)\s*\d", text):
        return True
    return False


def calculate_total_watts(unit_watts: object, quantity: object) -> Optional[str]:
    if has_multiple_wattage_options(unit_watts):
        return None
    watts_value = parse_decimal_from_text(unit_watts)
    quantity_value = parse_decimal_from_text(quantity)
    if watts_value is None or quantity_value is None:
        return None
    return format_watts_total(watts_value * quantity_value)


def extract_cct_options(value: object) -> Optional[str]:
    text = clean_optional_text(value)
    if not text:
        return None
    matches = re.findall(r"\b\d{3,5}\s*K\b", text, flags=re.I)
    cleaned = [re.sub(r"\s+", "", match).upper() for match in matches]
    return " / ".join(dict.fromkeys(cleaned)) if cleaned else None


def normalize_cct_value(value: object) -> Optional[str]:
    text = clean_optional_text(value)
    if not text:
        return None
    return extract_cct_options(text) or text


def extract_ip_rating(value: object) -> Optional[str]:
    text = clean_optional_text(value)
    if not text:
        return None
    match = re.search(r"\bIP(?:\s+[A-Za-z]+)?\s*[:#-]?\s*([0-9]{2})\b", text, flags=re.I)
    return f"IP{match.group(1)}" if match else None


def normalize_ip_rating(value: object) -> Optional[str]:
    text = clean_optional_text(value)
    if not text:
        return None
    extracted = extract_ip_rating(text)
    if extracted:
        return extracted
    if text.isdigit():
        return f"IP{text}"
    return text


def extract_dimming_options(value: object) -> Optional[str]:
    text = clean_optional_text(value)
    if not text:
        return None
    options: list[str] = []
    if re.search(r"without\s+dimmer|non[-\s]?dimmable|no\s+dimmer", text, flags=re.I):
        options.append("Without Dimmer")
    if re.search(r"0\s*[-–—]?\s*10\s*V", text, flags=re.I):
        options.append("0-10V")
    if re.search(r"\bDALI\b", text, flags=re.I):
        options.append("DALI")
    if re.search(r"\bCASAMBI\b", text, flags=re.I):
        options.append("CASAMBI")
    return " / ".join(dict.fromkeys(options)) if options else None


def normalize_dimming_value(value: object) -> Optional[str]:
    text = clean_optional_text(value)
    if not text:
        return None
    return extract_dimming_options(text) or text


def extract_lumen_options(value: object) -> Optional[str]:
    text = clean_optional_text(value)
    if not text:
        return None
    matches = re.findall(r"\b\d+(?:[.,]\d+)?\s*lm\b", text, flags=re.I)
    cleaned = [re.sub(r"\s+", "", match).replace(",", ".") for match in matches]
    return " / ".join(dict.fromkeys(cleaned)) if cleaned else None


def apply_wattage_multiplier_metadata(normalized: dict[str, Optional[str]]) -> None:
    unit_watts = clean_optional_text(normalized.get("unit_watts"))
    if not unit_watts:
        return
    match = re.search(r"(\d+(?:\.\d+)?)\s*W\s*(?:\*|x|×)\s*(\d+)", unit_watts, flags=re.I)
    if not match:
        return
    watts = Decimal(match.group(1))
    quantity = match.group(2)
    normalized["unit_watts"] = f"{format_watts_total(watts)}"
    normalized["quantity"] = normalized.get("quantity") or quantity
    calculated_total = calculate_total_watts(normalized.get("unit_watts"), normalized.get("quantity"))
    if calculated_total:
        normalized["total_watts"] = calculated_total


def apply_total_watts_calculation(row: Luminaire) -> None:
    total = calculate_total_watts(row.unit_watts, row.quantity)
    if total is not None:
        row.total_watts = total


def normalize_import_metadata(metadata: object) -> dict[str, Optional[str]]:
    if not isinstance(metadata, dict):
        return {}

    normalized_lookup = {
        canonicalize_import_key(key): clean_optional_text(value)
        for key, value in metadata.items()
    }

    normalized: dict[str, Optional[str]] = {}
    for target_key, aliases in IMPORT_METADATA_ALIASES.items():
        value = None
        for alias in aliases:
            alias_value = normalized_lookup.get(canonicalize_import_key(alias))
            if alias_value:
                value = alias_value
                break
        normalized[target_key] = value

    normalized["manufacturer"] = normalize_manufacturer_name(normalized.get("manufacturer"))
    normalized["cct"] = clean_optional_text(normalized.get("cct"))
    normalized["dimming"] = clean_optional_text(normalized.get("dimming"))
    normalized["ip_rating"] = clean_optional_text(normalized.get("ip_rating"))
    apply_wattage_multiplier_metadata(normalized)
    normalized["type"] = normalize_type_name(
        normalized.get("type"),
        product=normalized.get("product"),
        description=normalized.get("description"),
    )
    calculated_total = calculate_total_watts(normalized.get("unit_watts"), normalized.get("quantity"))
    if calculated_total is not None:
        normalized["total_watts"] = calculated_total
    return normalized


EMPTY_VALUE_TEXT = "N/A"
NA_METADATA_KEYS = (
    "type", "manufacturer", "product", "spec_code", "description", "colour", "cct", "cri",
    "unit_watts", "dimming", "quantity", "ip_rating", "total_watts",
)


def fill_empty_import_values(metadata: dict, custom_values: dict, custom_columns: list[dict]) -> None:
    """Extraction found nothing for these fields: show "N/A" instead of a blank cell."""
    for key in NA_METADATA_KEYS:
        if not clean_optional_text(metadata.get(key)):
            metadata[key] = EMPTY_VALUE_TEXT
    for column in custom_columns:
        name = column.get("name")
        if name and not clean_optional_text(custom_values.get(name)):
            custom_values[name] = EMPTY_VALUE_TEXT


def find_luminaire_reference_column_name(columns: list[dict]) -> Optional[str]:
    for column in columns:
        name = clean_optional_text(column.get("name"))
        if not name:
            continue
        if canonicalize_import_key(name) in LUMINAIRE_REFERENCE_COLUMN_ALIASES:
            return name
    return None


def extract_luminaire_reference_from_source_name(source_name: object) -> Optional[str]:
    raw_name = clean_optional_text(source_name)
    if not raw_name:
        return None

    parsed = urlparse(raw_name)
    candidate = parsed.path or raw_name
    candidate = unquote(candidate.rsplit("/", 1)[-1]).strip()
    stem = Path(candidate).stem.strip()
    return stem or None


def sanitize_export_filename_part(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9]+", "-", value).strip("-").lower()
    return cleaned or "luminaire-schedule"


def default_schedule_label(column_key: str) -> str:
    if column_key.startswith("custom:"):
        return "Custom Field"
    if column_key in SCHEDULE_EXPORT_LABELS:
        return SCHEDULE_EXPORT_LABELS[column_key]
    pretty = re.sub(r"([a-z0-9])([A-Z])", r"\1 \2", column_key.replace("_", " "))
    return pretty.strip().title() or "Column"


def schedule_column_width(column_key: str) -> float:
    widths = {
        "type": 78,
        "photo": 58,
        "diagram": 58,
        "polar": 58,
        "manufacturer": 100,
        "product": 112,
        "specCode": 92,
        "description": 170,
        "colour": 76,
        "cct": 62,
        "cri": 56,
        "unitWatts": 74,
        "dimming": 82,
        "quantity": 52,
        "ipRating": 68,
        "totalWatts": 82,
        "document": 110,
        "templatePdf": 110,
    }
    return widths.get(column_key, 96)


def get_schedule_export_image_path(row: Luminaire, column_key: str) -> Optional[Path]:
    relative_path: Optional[str] = None
    if column_key == "photo":
        relative_path = row.product_image_path
    elif column_key == "diagram":
        relative_path = row.diagram_image_path
    elif column_key == "polar":
        relative_path = row.polar_image_path

    if not relative_path:
        return None

    path = DATA_DIR / relative_path
    return path if path.exists() else None


def get_schedule_export_text(
    row: Luminaire,
    column_key: str,
    custom_values: dict,
    custom_columns_by_id: dict[int, CustomColumn],
    template_names_by_id: dict[int, str],
) -> str:
    if column_key.startswith("custom:"):
        try:
            custom_id = int(column_key.split(":", 1)[1])
        except Exception:
            return ""
        custom_column = custom_columns_by_id.get(custom_id)
        if not custom_column:
            return ""
        return str(custom_values.get(custom_column.name, "") or "")

    if column_key == "type":
        return row.type_code or ""
    if column_key == "photo":
        return to_public_file_url(row.product_image_path) or ""
    if column_key == "diagram":
        return to_public_file_url(row.diagram_image_path) or ""
    if column_key == "polar":
        return to_public_file_url(row.polar_image_path) or ""
    if column_key == "manufacturer":
        return normalize_manufacturer_name(row.manufacturer) or ""
    if column_key == "product":
        return row.product_name or ""
    if column_key == "specCode":
        return row.spec_code or ""
    if column_key == "description":
        return row.description or ""
    if column_key == "colour":
        return row.colour or ""
    if column_key == "cct":
        return row.cct or ""
    if column_key == "cri":
        return row.cri or ""
    if column_key == "unitWatts":
        return row.unit_watts or ""
    if column_key == "dimming":
        return row.dimming or ""
    if column_key == "quantity":
        return row.quantity or ""
    if column_key == "ipRating":
        return row.ip_rating or ""
    if column_key == "totalWatts":
        return row.total_watts or ""
    if column_key == "document":
        return row.source_url or to_public_file_url(row.document_path) or ""
    if column_key == "templatePdf":
        if row.pdf_template and str(row.pdf_template).isdigit():
            return template_names_by_id.get(int(str(row.pdf_template)), str(row.pdf_template))
        return row.pdf_template or ""
    return ""


def build_schedule_export_rows(
    luminaires: List[Luminaire],
    column_keys: List[str],
    column_labels: dict[str, str],
    custom_columns_by_id: dict[int, CustomColumn],
    template_names_by_id: dict[int, str],
) -> list[dict]:
    rows: list[dict] = []
    for luminaire in luminaires:
        custom_values = {}
        if luminaire.custom_data:
            try:
                custom_values = json.loads(luminaire.custom_data)
            except Exception:
                custom_values = {}

        cells = []
        for column_key in column_keys:
            text_value = get_schedule_export_text(
                luminaire,
                column_key,
                custom_values,
                custom_columns_by_id,
                template_names_by_id,
            )
            image_path = get_schedule_export_image_path(luminaire, column_key)
            label = column_labels.get(column_key)
            if not label:
                if column_key.startswith("custom:"):
                    try:
                        custom_id = int(column_key.split(":", 1)[1])
                    except Exception:
                        custom_id = 0
                    custom_column = custom_columns_by_id.get(custom_id)
                    label = custom_column.name if custom_column else default_schedule_label(column_key)
                else:
                    label = default_schedule_label(column_key)

            cells.append({
                "key": column_key,
                "label": label,
                "text": text_value.strip() or "—",
                "image_path": image_path,
            })

        rows.append({"id": luminaire.id, "cells": cells})

    return rows


def build_schedule_export_title(project_name: str) -> str:
    return f"{(project_name or '').strip() or 'Project'} - Luminaire Schedule"


def get_schedule_brand_logo_asset_path(source_path: Optional[Path] = None) -> Optional[Path]:
    source_path = source_path or SCHEDULE_BRAND_LOGO_SOURCE_PATH
    if not source_path.exists():
        return None
    if source_path == SCHEDULE_BRAND_LOGO_SOURCE_PATH:
        export_path = SCHEDULE_BRAND_LOGO_EXPORT_PATH
    else:
        export_path = BRANDING_DIR / f"schedule-{source_path.stem}.png"
        try:
            if export_path.exists() and export_path.stat().st_mtime >= source_path.stat().st_mtime:
                return export_path
        except OSError:
            pass

    try:
        from PIL import Image, ImageChops
    except ImportError:
        return source_path

    try:
        with Image.open(source_path) as source_image:
            source_image.load()
            logo_image = source_image.convert("RGBA")
            flattened_logo = Image.alpha_composite(
                Image.new("RGBA", logo_image.size, (255, 255, 255, 255)),
                logo_image,
            )
            flattened_rgb = flattened_logo.convert("RGB")
            background_color = flattened_rgb.getpixel((0, 0))
            background = Image.new("RGB", flattened_rgb.size, background_color)

            # Trim only the uniform outer background so the full logo stays readable.
            diff_mask = ImageChops.difference(flattened_rgb, background).convert("L")
            bbox = diff_mask.point(lambda value: 255 if value > 12 else 0).getbbox()
            if bbox:
                content_width = bbox[2] - bbox[0]
                content_height = bbox[3] - bbox[1]
                pad_x = max(10, int(content_width * 0.05))
                pad_y = max(10, int(content_height * 0.08))
                crop_box = (
                    max(0, bbox[0] - pad_x),
                    max(0, bbox[1] - pad_y),
                    min(flattened_logo.width, bbox[2] + pad_x),
                    min(flattened_logo.height, bbox[3] + pad_y),
                )
                prepared_logo = flattened_logo.crop(crop_box)
            else:
                prepared_logo = flattened_logo
            prepared_logo.save(export_path)
    except Exception:
        return source_path

    return export_path


def get_schedule_brand(db: Session, project_id: int) -> dict:
    """Logo / studio name / accent used on schedule exports for one project.

    Uses the project's own logo when set. Projects that hide the legacy default
    logo (client projects without a logo) get no logo, never the Afterglow one.
    """
    logo_source: Optional[Path] = None
    configured_path = get_setting(db, "specification_logo_path", "", project_id=project_id)
    if configured_path and (DATA_DIR / configured_path).exists():
        logo_source = DATA_DIR / configured_path
    elif get_setting(db, "brand_hide_default_logo", "", project_id=project_id) != "1":
        logo_source = SCHEDULE_BRAND_LOGO_SOURCE_PATH
    accent = normalize_accent(get_setting(db, "brand_accent", "", project_id=project_id))
    return {
        "logo_source": logo_source,
        "use_logo": logo_source is not None,
        "studio_name": (get_setting(db, "brand_studio_name", "", project_id=project_id) or "").strip(),
        "accent": accent or "",
    }


def schedule_header_hex(brand: Optional[dict]) -> str:
    accent = (brand or {}).get("accent") or ""
    return accent.lstrip("#").upper() if accent else "0B1C31"


def get_schedule_brand_logo(
    max_width: float,
    max_height: float,
    brand: Optional[dict] = None,
) -> Optional[tuple[Path, int, int]]:
    if brand is not None and not brand.get("use_logo"):
        return None
    logo_path = get_schedule_brand_logo_asset_path((brand or {}).get("logo_source"))
    if not logo_path:
        return None

    try:
        source_width, source_height = ImageReader(str(logo_path)).getSize()
    except Exception:
        return None

    if source_width <= 0 or source_height <= 0:
        return None

    scale = min(max_width / source_width, max_height / source_height)
    width = max(1, int(source_width * scale))
    height = max(1, int(source_height * scale))
    return logo_path, width, height


def build_schedule_excel_export(path: Path, project_name: str, rows: list[dict], brand: Optional[dict] = None) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise HTTPException(
            status_code=500,
            detail="Excel export requires openpyxl and Pillow to be installed on the backend.",
        ) from exc

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Schedule"
    sheet.sheet_view.showGridLines = False

    header_fill = PatternFill(fill_type="solid", fgColor=schedule_header_hex(brand))
    header_font = Font(color="FFFFFF", bold=True)
    body_alignment = Alignment(vertical="top", wrap_text=True)
    thin_side = Side(style="thin", color="D5DEEA")
    border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    headers = rows[0]["cells"] if rows else []
    logo = get_schedule_brand_logo(max_width=128, max_height=56, brand=brand)
    title = build_schedule_export_title(project_name)
    studio_name = (brand or {}).get("studio_name") or ""
    header_row_index = 3
    data_start_row_index = header_row_index + 1

    if logo:
        logo_path, logo_width, logo_height = logo
        try:
            image = XLImage(str(logo_path))
            image.width = logo_width
            image.height = logo_height
            sheet.add_image(image, "A1")
        except Exception:
            pass

    title_column_index = 2 if logo else 1
    title_end_column_index = max(len(headers), title_column_index)
    title_start_col = get_column_letter(title_column_index)
    title_end_col = get_column_letter(title_end_column_index)
    if title_end_column_index > title_column_index:
        sheet.merge_cells(f"{title_start_col}1:{title_end_col}1")

    title_cell = sheet.cell(row=1, column=title_column_index, value=title)
    title_cell.font = Font(color="0F172A", bold=True, size=18)
    title_cell.alignment = Alignment(vertical="center")
    if studio_name:
        studio_cell = sheet.cell(row=2, column=title_column_index, value=studio_name)
        studio_cell.font = Font(color="475569", size=10)

    sheet.row_dimensions[1].height = max(56, (logo[2] + 8) if logo else 56)
    sheet.row_dimensions[2].height = 16 if studio_name else 12
    sheet.freeze_panes = f"A{data_start_row_index}"

    for column_index, cell in enumerate(headers, start=1):
        header_cell = sheet.cell(row=header_row_index, column=column_index, value=cell["label"])
        header_cell.fill = header_fill
        header_cell.font = header_font
        header_cell.alignment = Alignment(horizontal="center", vertical="center")
        header_cell.border = border
        sheet.column_dimensions[get_column_letter(column_index)].width = max(schedule_column_width(cell["key"]) / 7.2, 10)

    if logo:
        sheet.column_dimensions["A"].width = max(sheet.column_dimensions["A"].width or 0, 14)

    sheet.row_dimensions[header_row_index].height = 24

    for row_index, row_data in enumerate(rows, start=data_start_row_index):
        row_has_image = False
        for column_index, cell in enumerate(row_data["cells"], start=1):
            target_cell = sheet.cell(row=row_index, column=column_index)
            target_cell.alignment = body_alignment
            target_cell.border = border

            image_path = cell["image_path"]
            if image_path:
                row_has_image = True
                try:
                    image = XLImage(str(image_path))
                    image.width = 48
                    image.height = 48
                    sheet.add_image(image, target_cell.coordinate)
                except Exception:
                    target_cell.value = cell["text"]
            else:
                target_cell.value = cell["text"]

        sheet.row_dimensions[row_index].height = 42 if row_has_image else 24

    if headers:
        last_col = get_column_letter(len(headers))
        sheet.auto_filter.ref = f"A{header_row_index}:{last_col}{len(rows) + header_row_index}"

    workbook.save(path)


def split_schedule_column_groups(widths: list[float], available: float, minimums: list[float]) -> list[list[int]]:
    """Split column indexes into page-width groups so no column is clipped."""
    groups: list[list[int]] = []
    current: list[int] = []
    used = 0.0
    for index, minimum in enumerate(minimums):
        if current and used + minimum > available:
            groups.append(current)
            current, used = [], 0.0
        current.append(index)
        used += minimum
    if current:
        groups.append(current)
    return groups or [[]]


def build_schedule_pdf_export(path: Path, project_name: str, rows: list[dict], brand: Optional[dict] = None) -> None:
    page_size = landscape(A3)
    document = SimpleDocTemplate(
        str(path),
        pagesize=page_size,
        leftMargin=18,
        rightMargin=18,
        topMargin=24,
        bottomMargin=24,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ScheduleExportTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        textColor=colors.HexColor("#0F172A"),
        spaceAfter=4,
    )
    header_style = ParagraphStyle(
        "ScheduleExportHeader",
        parent=styles["BodyText"],
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=10,
        textColor=colors.white,
    )
    body_style = ParagraphStyle(
        "ScheduleExportBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#0F172A"),
        wordWrap="CJK",
    )

    title_text = escape(build_schedule_export_title(project_name))
    studio_name = (brand or {}).get("studio_name") or ""
    if studio_name:
        title_text += f'<br/><font size="9" color="#475569">{escape(studio_name)}</font>'
    header_color = colors.HexColor("#" + schedule_header_hex(brand))
    logo = get_schedule_brand_logo(max_width=128, max_height=56, brand=brand)
    elements = []
    if logo:
        logo_path, logo_width, logo_height = logo
        header_table = Table(
            [[RLImage(str(logo_path), width=logo_width, height=logo_height), Paragraph(title_text, title_style)]],
            colWidths=[logo_width + 12, max(document.width - logo_width - 12, 120)],
        )
        header_table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 0),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ]))
        elements.extend([header_table, Spacer(1, 12)])
    else:
        elements.extend([Paragraph(title_text, title_style), Spacer(1, 12)])

    headers = rows[0]["cells"] if rows else []
    image_keys = {"photo", "diagram", "polar"}
    raw_widths = [schedule_column_width(cell["key"]) for cell in headers]
    minimums = [54 if cell["key"] in image_keys else 64 for cell in headers]
    available_width = document.width
    # Wide schedules are split into column groups (each a full-width table with a
    # row-number column) instead of being clipped at the page edge.
    ref_width = 26
    if sum(raw_widths) <= available_width:
        groups = [list(range(len(headers)))]
    elif sum(minimums) <= available_width:
        groups = [list(range(len(headers)))]
    else:
        groups = split_schedule_column_groups(raw_widths, available_width - ref_width, minimums)
    multi_group = len(groups) > 1

    def fit_widths(indexes: list[int], room: float) -> list[float]:
        widths = [raw_widths[i] for i in indexes]
        total = sum(widths) or 1
        if total > room:
            mins = [minimums[i] for i in indexes]
            spare = max(room - sum(mins), 0)
            extra = [max(w - m, 0) for w, m in zip(widths, mins)]
            extra_total = sum(extra) or 1
            return [m + spare * e / extra_total for m, e in zip(mins, extra)]
        return widths

    def cell_flowable(cell: dict):
        image_path = cell["image_path"]
        if image_path:
            try:
                return RLImage(str(image_path), width=42, height=42)
            except Exception:
                pass
        return Paragraph(escape(cell["text"]).replace("\n", "<br/>"), body_style)

    for group_number, indexes in enumerate(groups, start=1):
        if multi_group:
            caption = f"Part {group_number} of {len(groups)} \u00b7 {headers[indexes[0]]['label']} \u2192 {headers[indexes[-1]]['label']}"
            elements.append(Paragraph(escape(caption), ParagraphStyle(
                "ScheduleGroupCaption", parent=body_style, fontName="Helvetica-Bold", fontSize=9, leading=12,
                textColor=colors.HexColor("#475569"), spaceAfter=4,
            )))
            column_widths = [ref_width] + fit_widths(indexes, available_width - ref_width)
            table_rows = [[Paragraph("#", header_style)] + [Paragraph(escape(headers[i]["label"]), header_style) for i in indexes]]
            for row_number, row in enumerate(rows, start=1):
                table_rows.append([Paragraph(str(row_number), body_style)] + [cell_flowable(row["cells"][i]) for i in indexes])
        else:
            column_widths = fit_widths(indexes, available_width)
            table_rows = [[Paragraph(escape(headers[i]["label"]), header_style) for i in indexes]]
            for row in rows:
                table_rows.append([cell_flowable(row["cells"][i]) for i in indexes])
        elements.append(_styled_schedule_table(table_rows, column_widths, header_color))
        if multi_group and group_number < len(groups):
            elements.append(PageBreak())
    document.build(elements)


def _styled_schedule_table(table_rows: list, column_widths: list[float], header_color) -> "Table":
    schedule_table = Table(table_rows, colWidths=column_widths, repeatRows=1)
    schedule_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), header_color),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (0, 0), (-1, 0), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FBFF")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    return schedule_table


def save_upload(file: UploadFile) -> tuple[Path, str]:
    extension = Path(file.filename or "document.pdf").suffix or ".pdf"
    unique_name = f"{uuid.uuid4().hex}{extension}"
    target_path = DOCUMENTS_DIR / unique_name
    with target_path.open("wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
    return target_path, f"documents/{unique_name}"


def save_pdf_bytes(content: bytes, url: str = "") -> tuple[Path, str]:
    unique_name = f"{uuid.uuid4().hex}.pdf"
    target_path = DOCUMENTS_DIR / unique_name
    target_path.write_bytes(content)
    return target_path, f"documents/{unique_name}"


def fetch_url_document(url: str) -> dict:
    """Download a submitted URL. Returns {'kind': 'pdf', ...} or {'kind': 'html', ...}."""
    response, content, final_url = safe_get(url, max_bytes=max(MAX_PDF_BYTES, MAX_PAGE_BYTES))
    content_type = response.headers.get("content-type", "")
    if looks_like_pdf(content, content_type, final_url):
        saved_path, stored_document = save_pdf_bytes(content, final_url)
        disposition = response.headers.get("content-disposition", "")
        match = re.search(r"filename\*=(?:UTF-8'')?([^;]+)|filename=\"?([^\";]+)\"?", disposition, flags=re.I)
        filename = unquote((match.group(1) or match.group(2)).strip().strip('"')) if match else ""
        return {
            "kind": "pdf",
            "path": saved_path,
            "stored_document": stored_document,
            "final_url": final_url,
            "filename": Path(filename).name if filename else "",
        }
    if "html" in content_type.lower() or content.lstrip()[:15].lower().startswith((b"<!doctype", b"<html")):
        if len(content) > MAX_PAGE_BYTES:
            raise ValueError(f"Web page is too large to read: {url}")
        return {"kind": "html", "html": decode_html(response, content), "final_url": final_url}
    raise ValueError(f"URL is neither a PDF nor a product web page: {url}")


IMPORT_WORKERS = max(1, int(os.getenv("IMPORT_PARALLEL_WORKERS", "3")))


def process_import_item(item: tuple[str, object], api_key: str, custom_columns: list[dict]) -> dict:
    """Process one uploaded file or URL. Runs in a worker thread."""
    kind, payload = item
    if kind == "upload":
        upload: UploadFile = payload  # type: ignore[assignment]
        source = upload.filename or "upload"
        if not source.lower().endswith(".pdf"):
            return {"ok": False, "chargeable": False, "source": source, "error": "Only PDF uploads are supported."}
        saved_path = None
        try:
            saved_path, stored_document = save_upload(upload)
            result = process_luminaire_document(str(saved_path), api_key=api_key, custom_columns=custom_columns)
            return {"ok": True, "row": ("upload", source, None, stored_document, result)}
        except Exception as exc:
            if saved_path and saved_path.exists():
                saved_path.unlink(missing_ok=True)
            return {"ok": False, "chargeable": True, "source": source, "error": str(exc)}

    url = str(payload)
    saved_path = None
    try:
        fetched = fetch_url_document(url)
        if fetched["kind"] == "pdf":
            saved_path = fetched["path"]
            result = process_luminaire_document(str(saved_path), api_key=api_key, custom_columns=custom_columns)
            source_name = (
                fetched.get("filename")
                or Path(urlparse(fetched["final_url"]).path).name
                or Path(urlparse(url).path).name
                or url
            )
            return {"ok": True, "row": ("url", source_name, url, fetched["stored_document"], result)}

        result = process_luminaire_webpage(
            url,
            fetched["html"],
            fetched["final_url"],
            api_key,
            custom_columns,
            save_pdf_bytes,
        )
        source_type = result.pop("source_type", "webpage")
        source_name = result.pop("source_name", None) or url
        stored_document = result.pop("stored_document", "") or ""
        source_url = result.pop("source_document_url", None) or url
        return {"ok": True, "row": (source_type, source_name, source_url, stored_document, result)}
    except Exception as exc:
        if saved_path and Path(saved_path).exists():
            Path(saved_path).unlink(missing_ok=True)
        return {"ok": False, "chargeable": True, "source": url, "error": str(exc)}


def process_import_items(items: list[tuple[str, object]], api_key: str, custom_columns: list[dict]) -> list[dict]:
    if len(items) <= 1:
        return [process_import_item(item, api_key, custom_columns) for item in items]
    with ThreadPoolExecutor(max_workers=min(IMPORT_WORKERS, len(items))) as executor:
        return list(executor.map(lambda item: process_import_item(item, api_key, custom_columns), items))


def refund_freemium_credits(user_id: int, credits: int) -> Optional[int]:
    """Give back credits reserved for documents that did not produce a row."""
    if credits <= 0:
        return None
    with session_scope() as db:
        db.query(User).filter(
            User.id == user_id,
            User.account_type == "freemium",
        ).update({User.tokens: User.tokens + credits}, synchronize_session=False)
        db.commit()
        return db.query(User.tokens).filter(User.id == user_id).scalar() or 0


def save_extracted_asset(temp_path: Optional[str], destination_dir: Path, folder_name: str) -> Optional[str]:
    if not temp_path:
        return None
    source_path = Path(temp_path)
    if not source_path.exists():
        return None
    new_name = f"{uuid.uuid4().hex}{source_path.suffix.lower() or '.png'}"
    final_path = destination_dir / new_name
    shutil.move(str(source_path), final_path)
    return f"images/{folder_name}/{new_name}"


@app.get("/")
async def index():
    return {"status": "ok", "message": "Luminaire API is running."}




def normalize_column_keys(values: List[str]) -> List[str]:
    normalized: List[str] = []
    seen = set()
    for item in values:
        key = str(item).strip()
        if not key or key in seen:
            continue
        seen.add(key)
        normalized.append(key)
    return normalized


def get_saved_column_order(db: Session, project_id: int) -> List[str]:
    raw = get_setting(db, "luminaire_column_order", "[]", project_id=project_id)
    try:
        parsed = json.loads(raw)
    except Exception:
        return []
    if not isinstance(parsed, list):
        return []
    return normalize_column_keys([str(item) for item in parsed])


def save_column_order_setting(db: Session, order: List[str], project_id: int) -> List[str]:
    normalized = normalize_column_keys(order)
    set_setting(db, "luminaire_column_order", json.dumps(normalized), project_id=project_id)
    return normalized


def get_hidden_columns(db: Session, project_id: int) -> List[str]:
    raw = get_setting(db, "luminaire_hidden_columns", "[]", project_id=project_id)
    try:
        parsed = json.loads(raw)
    except Exception:
        return []
    if not isinstance(parsed, list):
        return []
    return normalize_column_keys([str(item) for item in parsed])


def save_hidden_columns_setting(db: Session, hidden: List[str], project_id: int) -> List[str]:
    normalized = normalize_column_keys(hidden)
    set_setting(db, "luminaire_hidden_columns", json.dumps(normalized), project_id=project_id)
    return normalized


def get_sort_config(db: Session, project_id: int) -> List[dict]:
    raw = get_setting(db, "luminaire_sort_config", "[]", project_id=project_id)
    try:
        parsed = json.loads(raw)
    except Exception:
        return []
    if not isinstance(parsed, list):
        return []

    normalized = []
    for item in parsed:
        if not isinstance(item, dict):
            continue
        key = str(item.get("key", "")).strip()
        direction = str(item.get("direction", "asc")).strip().lower()
        if not key:
            continue
        if direction not in {"asc", "desc"}:
            direction = "asc"
        normalized.append({"key": key, "direction": direction})
    return normalized[:1]


def save_sort_config_setting(db: Session, sort_config: List[dict], project_id: int) -> List[dict]:
    normalized = []
    for item in sort_config:
        if not isinstance(item, dict):
            continue
        key = str(item.get("key", "")).strip()
        direction = str(item.get("direction", "asc")).strip().lower()
        if not key:
            continue
        if direction not in {"asc", "desc"}:
            direction = "asc"
        normalized.append({"key": key, "direction": direction})
    normalized = normalized[:1]
    set_setting(db, "luminaire_sort_config", json.dumps(normalized), project_id=project_id)
    return normalized


def get_filter_config(db: Session, project_id: int) -> List[dict]:
    raw = get_setting(db, "luminaire_filter_config", "[]", project_id=project_id)
    try:
        parsed = json.loads(raw)
    except Exception:
        return []
    if not isinstance(parsed, list):
        return []

    allowed_ops = {"contains", "not_contains", "is", "is_not", "is_empty", "is_not_empty", "gt", "lt"}
    normalized = []
    for item in parsed:
        if not isinstance(item, dict):
            continue
        key = str(item.get("key", "")).strip()
        operator = str(item.get("operator", "contains")).strip().lower()
        if not key:
            continue
        if operator not in allowed_ops:
            operator = "contains"
        normalized.append({
            "id": str(item.get("id", uuid.uuid4().hex)),
            "key": key,
            "operator": operator,
            "value": str(item.get("value", "")),
        })
    return normalized


def save_filter_config_setting(db: Session, filters: List[dict], project_id: int) -> List[dict]:
    allowed_ops = {"contains", "not_contains", "is", "is_not", "is_empty", "is_not_empty", "gt", "lt"}
    normalized = []
    for item in filters:
        if not isinstance(item, dict):
            continue
        key = str(item.get("key", "")).strip()
        operator = str(item.get("operator", "contains")).strip().lower()
        if not key:
            continue
        if operator not in allowed_ops:
            operator = "contains"
        normalized.append({
            "id": str(item.get("id", uuid.uuid4().hex)),
            "key": key,
            "operator": operator,
            "value": str(item.get("value", "")),
        })
    set_setting(db, "luminaire_filter_config", json.dumps(normalized), project_id=project_id)
    return normalized


def get_column_labels(db: Session, project_id: int) -> dict:
    raw = get_setting(db, "luminaire_column_labels", "{}", project_id=project_id)
    try:
        parsed = json.loads(raw)
    except Exception:
        return {}
    if not isinstance(parsed, dict):
        return {}
    normalized = {}
    for key, value in parsed.items():
        safe_key = str(key).strip()
        safe_value = str(value).strip()
        if safe_key and safe_value:
            normalized[safe_key] = safe_value
    return normalized


def save_column_labels_setting(db: Session, labels: dict, project_id: int) -> dict:
    normalized = {}
    for key, value in labels.items():
        safe_key = str(key).strip()
        safe_value = str(value).strip()
        if safe_key and safe_value:
            normalized[safe_key] = safe_value
    set_setting(db, "luminaire_column_labels", json.dumps(normalized), project_id=project_id)
    return normalized


def serialize_custom_column(row: CustomColumn) -> dict:
    return {
        "id": row.id,
        "name": row.name,
        "aiEnabled": bool(row.ai_enabled),
        "aiDescription": row.ai_description or "",
    }


@app.get("/api/projects")
def list_projects(current_user: User = Depends(get_current_user)):
    with session_scope() as db:
        if current_user.role == "admin":
            rows = db.query(Project).order_by(Project.id.asc()).all()
        else:
            rows = db.query(Project).filter(Project.owner_id == current_user.id).order_by(Project.id.asc()).all()
        counts = dict(
            db.query(Luminaire.project_id, func.count(Luminaire.id))
            .filter(Luminaire.project_id.in_([r.id for r in rows] or [0]))
            .group_by(Luminaire.project_id)
            .all()
        )
        doc_counts = dict(
            db.query(Luminaire.project_id, func.count(GeneratedDocument.id))
            .join(GeneratedDocument, GeneratedDocument.luminaire_id == Luminaire.id)
            .filter(Luminaire.project_id.in_([r.id for r in rows] or [0]))
            .group_by(Luminaire.project_id)
            .all()
        )
        ids = [r.id for r in rows] or [0]
        meta: dict[tuple[int, str], str] = {}
        for setting in db.query(AppSetting).filter(
            AppSetting.project_id.in_(ids),
            AppSetting.key.in_(["specification_location_country", "specification_revision", "brand_template", "brand_studio_name", "schedule_exports"]),
        ).all():
            meta[(setting.project_id, setting.key)] = setting.value or ""

        def schedule_count(pid: int) -> int:
            try:
                return len(json.loads(meta.get((pid, "schedule_exports"), "") or "[]"))
            except Exception:
                return 0

        return {
            "projects": [
                {
                    **serialize_project(r),
                    "luminaireCount": int(counts.get(r.id, 0)),
                    "documentCount": int(doc_counts.get(r.id, 0)) + schedule_count(r.id),
                    "location": meta.get((r.id, "specification_location_country"), ""),
                    "revision": meta.get((r.id, "specification_revision"), "") or "A",
                    "template": normalize_brand_template(meta.get((r.id, "brand_template"), "standard")),
                    "studioName": meta.get((r.id, "brand_studio_name"), ""),
                }
                for r in rows
            ]
        }


@app.post("/api/projects")
async def create_project(request: Request, current_user: User = Depends(get_current_user)):
    data = await request.json()
    name = data.get("name", "").strip()
    location = str(data.get("location", "") or "").strip()[:100]
    revision = str(data.get("revision", "") or "").strip()[:20] or "A"
    issue_date = str(data.get("issueDate", "") or "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="Name required")
    if issue_date:
        try:
            datetime.strptime(issue_date, "%Y-%m-%d")
        except ValueError:
            raise HTTPException(status_code=400, detail="Issue date must be YYYY-MM-DD")
    with session_scope() as db:
        new_project = Project(name=name[:100], owner_id=current_user.id)
        db.add(new_project)
        db.flush()
        pid = new_project.id
        set_setting(db, "specification_location_country", location, project_id=pid)
        set_setting(db, "specification_revision", revision, project_id=pid)
        set_setting(db, "specification_issue_date", issue_date or datetime.utcnow().date().isoformat(), project_id=pid)
        # New projects inherit the owner's studio defaults (RISE OS Setup).
        owner = db.query(User).filter(User.id == current_user.id).first() or current_user
        studio = get_studio_defaults(db, owner)
        set_setting(db, "brand_studio_name", studio["studioName"], project_id=pid)
        set_setting(db, "brand_accent", studio["accent"], project_id=pid)
        set_setting(db, "brand_template", studio["template"], project_id=pid)
        if studio["logoPath"]:
            set_setting(db, "specification_logo_path", studio["logoPath"], project_id=pid)
        elif current_user.role != "admin":
            # Never show another studio's legacy logo on a client's sheets.
            set_setting(db, "brand_hide_default_logo", "1", project_id=pid)
        ensure_project_default_template(db, new_project.id)
        if studio["template"] == "compact":
            ensure_project_compact_template(db, pid)
        ensure_project_specification_columns(db, new_project.id)
        db.commit()
        db.refresh(new_project)
        return serialize_project(new_project)


@app.get("/api/studio")
def api_get_studio(current_user: User = Depends(get_current_user)):
    with session_scope() as db:
        user = db.query(User).filter(User.id == current_user.id).first()
        data = get_studio_defaults(db, user)
        data.pop("logoPath", None)
        return data


@app.put("/api/studio")
async def api_save_studio(request: Request, current_user: User = Depends(get_current_user)):
    data = await request.json()
    name = str(data.get("studioName", "") or "").strip()[:100]
    if not name:
        raise HTTPException(status_code=400, detail="Enter a studio name.")
    accent = normalize_accent(data.get("accent")) or DEFAULT_BRAND_ACCENT
    template = normalize_brand_template(data.get("template"))
    with session_scope() as db:
        user = db.query(User).filter(User.id == current_user.id).first()
        user.studio_name = name
        set_setting(db, studio_setting_key(user.id, "accent"), accent)
        set_setting(db, studio_setting_key(user.id, "template"), template)
        if data.get("removeLogo"):
            set_setting(db, studio_setting_key(user.id, "logo_path"), "")
        set_setting(db, studio_setting_key(user.id, "saved"), "1")
        db.commit()
        result = get_studio_defaults(db, user)
        result.pop("logoPath", None)
        return result


@app.post("/api/studio/logo")
async def api_upload_studio_logo(file: UploadFile = File(...), current_user: User = Depends(get_current_user)):
    relative_path = save_brand_logo(await file.read(), f"studio-{current_user.id}")
    with session_scope() as db:
        set_setting(db, studio_setting_key(current_user.id, "logo_path"), relative_path)
        db.commit()
    return {"logoUrl": to_public_file_url(relative_path)}


@app.get("/api/projects/{project_id}/brand")
def api_get_project_brand(project_id: int, current_user: User = Depends(get_current_user)):
    with session_scope() as db:
        project = get_project_or_404(db, project_id, current_user)
        return get_project_brand(db, project)


@app.put("/api/projects/{project_id}/brand")
async def api_save_project_brand(project_id: int, request: Request, current_user: User = Depends(get_current_user)):
    data = await request.json()
    name = str(data.get("name", "") or "").strip()[:100]
    studio_name = str(data.get("studioName", "") or "").strip()[:100]
    if not name or not studio_name:
        raise HTTPException(status_code=400, detail="Enter the project and studio names.")
    issue_date = str(data.get("issueDate", "") or "").strip()
    if issue_date:
        try:
            datetime.strptime(issue_date, "%Y-%m-%d")
        except ValueError:
            raise HTTPException(status_code=400, detail="Issue date must be YYYY-MM-DD")
    with session_scope() as db:
        project = get_project_or_404(db, project_id, current_user)
        project.name = name
        set_setting(db, "specification_location_country", str(data.get("location", "") or "").strip()[:100], project_id=project_id)
        set_setting(db, "specification_revision", str(data.get("revision", "") or "").strip()[:20] or "A", project_id=project_id)
        set_setting(db, "specification_issue_date", issue_date or datetime.utcnow().date().isoformat(), project_id=project_id)
        set_setting(db, "brand_studio_name", studio_name, project_id=project_id)
        set_setting(db, "brand_accent", normalize_accent(data.get("accent")) or DEFAULT_BRAND_ACCENT, project_id=project_id)
        if data.get("useStudioLogo"):
            user = db.query(User).filter(User.id == current_user.id).first()
            studio = get_studio_defaults(db, user)
            if studio["logoPath"]:
                set_setting(db, "specification_logo_path", studio["logoPath"], project_id=project_id)
        new_template = normalize_brand_template(data.get("template"))
        if new_template != normalize_brand_template(get_setting(db, "brand_template", "standard", project_id=project_id)):
            switch_project_brand_template(db, project_id, new_template)
        db.commit()
        db.refresh(project)
        return get_project_brand(db, project)


@app.put("/api/projects/{project_id}")
async def update_project(project_id: int, request: Request, current_user: User = Depends(get_current_user)):
    data = await request.json()
    name = data.get("name", "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="Name required")
    with session_scope() as db:
        project = get_project_or_404(db, project_id, current_user)
        project.name = name
        db.commit()
        db.refresh(project)
        return serialize_project(project)


@app.delete("/api/projects/{project_id}")
def delete_project(project_id: int, current_user: User = Depends(get_current_user)):
    with session_scope() as db:
        project = get_project_or_404(db, project_id, current_user)
        db.delete(project)
        db.commit()
    return {"status": "ok"}


@app.get("/api/columns")
def list_columns(project_id: int, current_user: User = Depends(get_current_user)):
    with session_scope() as db:
        get_project_or_404(db, project_id, current_user)
        ensure_project_specification_columns(db, project_id)
        db.commit()
        rows = db.query(CustomColumn).filter(CustomColumn.project_id == project_id).all()
        saved_order = get_saved_column_order(db, project_id)
        custom_positions = {key: idx for idx, key in enumerate(saved_order) if key.startswith("custom:")}
        rows.sort(key=lambda row: custom_positions.get(f"custom:{row.id}", len(saved_order) + row.id))
        return [serialize_custom_column(r) for r in rows]


@app.get("/api/columns/order")
def get_column_order(project_id: int, current_user: User = Depends(get_current_user)):
    with session_scope() as db:
        get_project_or_404(db, project_id, current_user)
        return {
            "order": get_saved_column_order(db, project_id),
            "hidden": get_hidden_columns(db, project_id),
            "sort": get_sort_config(db, project_id),
            "filter": get_filter_config(db, project_id),
            "labels": get_column_labels(db, project_id),
            "filterLogic": get_setting(db, "luminaire_filter_logic", "AND", project_id=project_id),
            "rowHeight": get_setting(db, "luminaire_row_height", "medium", project_id=project_id),
        }


@app.put("/api/columns/order")
async def update_column_order(request: Request, project_id: int, current_user: User = Depends(get_current_user)):
    data = await request.json()
    order = data.get("order", [])
    hidden = data.get("hidden", [])
    sort_config = data.get("sort", [])
    filter_config = data.get("filter", [])
    labels = data.get("labels", {})
    filter_logic = data.get("filterLogic", "AND")
    row_height = data.get("rowHeight", "medium")
    if filter_logic not in ("AND", "OR"):
        filter_logic = "AND"
    if row_height not in ("small", "medium", "large"):
        row_height = "medium"
    if not isinstance(order, list):
        raise HTTPException(status_code=400, detail="Order must be a list")
    if not isinstance(hidden, list):
        raise HTTPException(status_code=400, detail="Hidden must be a list")
    if not isinstance(sort_config, list):
        raise HTTPException(status_code=400, detail="Sort must be a list")
    if not isinstance(filter_config, list):
        raise HTTPException(status_code=400, detail="Filter must be a list")
    if not isinstance(labels, dict):
        raise HTTPException(status_code=400, detail="Labels must be an object")
    with session_scope() as db:
        get_project_or_404(db, project_id, current_user)
        normalized_order = save_column_order_setting(db, [str(item) for item in order], project_id)
        normalized_hidden = save_hidden_columns_setting(db, [str(item) for item in hidden], project_id)
        normalized_sort = save_sort_config_setting(db, sort_config, project_id)
        normalized_filter = save_filter_config_setting(db, filter_config, project_id)
        normalized_labels = save_column_labels_setting(db, labels, project_id)
        set_setting(db, "luminaire_filter_logic", filter_logic, project_id=project_id)
        set_setting(db, "luminaire_row_height", row_height, project_id=project_id)
        db.commit()
        return {
            "order": normalized_order,
            "hidden": normalized_hidden,
            "sort": normalized_sort,
            "filter": normalized_filter,
            "labels": normalized_labels,
            "filterLogic": filter_logic,
            "rowHeight": row_height,
        }


@app.post("/api/columns")
async def create_column(request: Request, project_id: int, current_user: User = Depends(get_current_user)):
    data = await request.json()
    name = data.get("name", "").strip()
    ai_enabled = bool(data.get("aiEnabled", False))
    ai_description = str(data.get("aiDescription", "") or "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="Name required")
    if is_standard_column_name(name):
        raise HTTPException(status_code=400, detail="This column already exists as a standard table column.")
    with session_scope() as db:
        get_project_or_404(db, project_id, current_user)
        exist = db.query(CustomColumn).filter(CustomColumn.name == name, CustomColumn.project_id == project_id).first()
        if exist:
            exist.ai_enabled = ai_enabled
            exist.ai_description = ai_description or None
            db.commit()
            db.refresh(exist)
            return serialize_custom_column(exist)
        new_col = CustomColumn(
            name=name,
            project_id=project_id,
            ai_enabled=ai_enabled,
            ai_description=ai_description or None,
        )
        db.add(new_col)
        db.flush()
        saved_order = get_saved_column_order(db, project_id)
        new_key = f"custom:{new_col.id}"
        if new_key not in saved_order:
            saved_order.append(new_key)
        save_column_order_setting(db, saved_order, project_id)
        db.commit()
        db.refresh(new_col)
        return serialize_custom_column(new_col)


@app.delete("/api/columns/{column_id}")
def delete_column(column_id: int, project_id: int, current_user: User = Depends(get_current_user)):
    with session_scope() as db:
        get_project_or_404(db, project_id, current_user)
        col = db.query(CustomColumn).filter(CustomColumn.id == column_id, CustomColumn.project_id == project_id).first()
        if not col:
            raise HTTPException(status_code=404, detail="Not found")
        db.delete(col)
        column_key = f"custom:{column_id}"
        save_column_order_setting(db, [key for key in get_saved_column_order(db, project_id) if key != column_key], project_id)
        save_hidden_columns_setting(db, [key for key in get_hidden_columns(db, project_id) if key != column_key], project_id)
        save_sort_config_setting(db, [item for item in get_sort_config(db, project_id) if item.get("key") != column_key], project_id)
        save_filter_config_setting(db, [item for item in get_filter_config(db, project_id) if item.get("key") != column_key], project_id)
        labels = get_column_labels(db, project_id)
        if column_key in labels:
            labels.pop(column_key, None)
            save_column_labels_setting(db, labels, project_id)
        db.commit()
    return {"status": "ok"}


@app.put("/api/columns/{column_id}")
async def update_column(column_id: int, request: Request, current_user: User = Depends(get_current_user)):
    data = await request.json()
    new_name = data.get("name", "").strip()
    ai_enabled = bool(data.get("aiEnabled", False))
    ai_description = str(data.get("aiDescription", "") or "").strip()
    if not new_name:
        raise HTTPException(status_code=400, detail="Name required")
    if is_standard_column_name(new_name):
        raise HTTPException(status_code=400, detail="This column already exists as a standard table column.")
    with session_scope() as db:
        col = db.query(CustomColumn).filter(CustomColumn.id == column_id).first()
        if not col:
            raise HTTPException(status_code=404, detail="Not found")
        get_project_or_404(db, col.project_id, current_user)

        old_name = col.name
        rows_to_update = db.query(Luminaire).filter(Luminaire.project_id == col.project_id).all()
        for r in rows_to_update:
            try:
                if r.custom_data:
                    custom = json.loads(r.custom_data)
                    if old_name in custom:
                        custom[new_name] = custom.pop(old_name)
                        r.custom_data = json.dumps(custom)
            except Exception:
                pass

        col.name = new_name
        col.ai_enabled = ai_enabled
        col.ai_description = ai_description or None
        db.commit()
        db.refresh(col)
        return serialize_custom_column(col)


@app.get("/api/luminaires")
def list_luminaires(project_id: int, current_user: User = Depends(get_current_user)):
    with session_scope() as db:
        get_project_or_404(db, project_id, current_user)
        default_template = project_brand_template(db, project_id)
        rows = db.query(Luminaire).filter(Luminaire.project_id == project_id).order_by(Luminaire.id.asc()).all()
        changed = False
        for row in rows:
            if not str(row.pdf_template or "").isdigit():
                row.pdf_template = str(default_template.id)
                changed = True
        if changed:
            db.commit()
        return {"items": [serialize_luminaire(row) for row in rows]}


@app.post("/api/luminaires")
def create_empty_luminaire(project_id: int, current_user: User = Depends(get_current_user)):
    with session_scope() as db:
        get_project_or_404(db, project_id, current_user)
        default_template = project_brand_template(db, project_id)
        new_row = Luminaire(
            project_id=project_id,
            source_type="manual",
            source_name="Manual Entry",
            source_url=None,
            document_path="",
            pdf_template=str(default_template.id),
        )
        db.add(new_row)
        db.commit()
        db.refresh(new_row)
        return serialize_luminaire(new_row)

@app.get("/api/luminaires/{item_id}")
def read_luminaire(item_id: int, current_user: User = Depends(get_current_user)):
    with session_scope() as db:
        row = db.query(Luminaire).filter(Luminaire.id == item_id).first()
        if not row:
            raise HTTPException(status_code=404, detail="Not found")
        get_project_or_404(db, row.project_id, current_user)
        return serialize_luminaire(row)


@app.get("/api/luminaires/{item_id}/designer-context")
def read_luminaire_designer_context(item_id: int, current_user: User = Depends(get_current_user)):
    with session_scope() as db:
        row = db.query(Luminaire).filter(Luminaire.id == item_id).first()
        if not row:
            raise HTTPException(status_code=404, detail="Not found")
        get_project_or_404(db, row.project_id, current_user)
        return designer_context_from_luminaire(db, row)


@app.delete("/api/luminaires/{item_id}")
def delete_luminaire(item_id: int, current_user: User = Depends(get_current_user)):
    with session_scope() as db:
        row = db.query(Luminaire).filter(Luminaire.id == item_id).first()
        if not row:
            raise HTTPException(status_code=404, detail="Not found")
        get_project_or_404(db, row.project_id, current_user)
        db.delete(row)
        db.commit()
    return {"status": "ok"}

@app.put("/api/luminaires/{luminaire_id}")
async def update_luminaire(luminaire_id: int, request: Request, current_user: User = Depends(get_current_user)):
    data = await request.json()
    with session_scope() as db:
        row = db.query(Luminaire).filter(Luminaire.id == luminaire_id).first()
        if not row:
            raise HTTPException(status_code=404, detail="Luminaire not found")
        get_project_or_404(db, row.project_id, current_user)
        
        # update fields
        if "manufacturer" in data: row.manufacturer = normalize_manufacturer_name(data["manufacturer"])
        if "product" in data: row.product_name = data["product"]
        if "specCode" in data: row.spec_code = data["specCode"]
        if "colour" in data: row.colour = data["colour"]
        if "cct" in data: row.cct = data["cct"]
        if "cri" in data: row.cri = data["cri"]
        if "ipRating" in data: row.ip_rating = data["ipRating"]
        if "unitWatts" in data: row.unit_watts = data["unitWatts"]
        if "dimming" in data: row.dimming = data["dimming"]
        if "type" in data: row.type_code = data["type"]
        if "description" in data: row.description = data["description"]
        if "quantity" in data: row.quantity = data["quantity"]
        if "unitWatts" in data or "quantity" in data:
            apply_total_watts_calculation(row)
        if "pdfTemplate" in data: row.pdf_template = data["pdfTemplate"]
        if "productImageUrl" in data: 
            val = data["productImageUrl"]
            row.product_image_path = val.replace("/files/", "") if val and val.startswith("/files/") else val
        if "diagramImageUrl" in data:
            val = data["diagramImageUrl"]
            row.diagram_image_path = val.replace("/files/", "") if val and val.startswith("/files/") else val
        if "polarImageUrl" in data:
            val = data["polarImageUrl"]
            row.polar_image_path = val.replace("/files/", "") if val and val.startswith("/files/") else val
        
        # update customData recursively
        if "customData" in data and isinstance(data["customData"], dict):
            existing_custom = {}
            if row.custom_data:
                try:
                    existing_custom = json.loads(row.custom_data)
                except Exception:
                    pass
            # merge new custom data completely
            row.custom_data = json.dumps({**existing_custom, **data["customData"]})
        
        db.commit()
        return serialize_luminaire(row)


@app.post("/api/luminaires/import")
async def import_luminaires(
    project_id: int,
    files: List[UploadFile] = File(default=[]),
    urls: str = Form(default="[]"),
    api_key: str = Form(default=""),
    current_user: User = Depends(get_current_user),
):
    submitted_urls = [item.strip() for item in json.loads(urls or "[]") if item.strip()]
    items_count = len(files) + len(submitted_urls)
    if items_count == 0:
        raise HTTPException(status_code=400, detail="Upload at least one PDF or add at least one URL.")
    if items_count > 20:
        raise HTTPException(status_code=400, detail="You can process at most 20 documents at once.")

    chargeable_items_count = (
        sum(1 for upload in files if (upload.filename or "").lower().endswith(".pdf"))
        + len(submitted_urls)
    )
    credits_remaining = None

    with session_scope() as db:
        get_project_or_404(db, project_id, current_user)
        ensure_project_specification_columns(db, project_id)
        db.commit()
        effective_api_key = (api_key or "").strip() or get_setting(db, "openai_api_key")
        project_custom_columns = db.query(CustomColumn).filter(CustomColumn.project_id == project_id).order_by(CustomColumn.id.asc()).all()
        hidden_column_keys = set(get_hidden_columns(db, project_id))
        saved_column_order = get_saved_column_order(db, project_id)
        column_order_positions = {
            key: index
            for index, key in enumerate(saved_column_order)
        }
        import_custom_columns = [
            {
                "name": col.name,
                "description": col.ai_description or "",
            }
            for col in project_custom_columns
            if not is_standard_column_name(col.name)
        ]
        ai_custom_columns = [
            {
                "name": col.name,
                "description": col.ai_description or "",
                "visible": f"custom:{col.id}" not in hidden_column_keys,
                "order": column_order_positions.get(
                    f"custom:{col.id}",
                    len(saved_column_order) + col.id,
                ),
            }
            for col in project_custom_columns
            if col.ai_enabled and not is_standard_column_name(col.name)
        ]
        ai_custom_columns.sort(
            key=lambda column: (
                not column["visible"],
                column["order"],
            )
        )
        if not effective_api_key:
            raise HTTPException(status_code=400, detail="Set the OpenAI API key in Configuration or provide it here for this import.")
        credits_remaining = reserve_freemium_credits(db, current_user.id, chargeable_items_count)

    created_rows = []
    failures = []
    import_items: list[tuple[str, object]] = [("upload", upload) for upload in files] + [
        ("url", url) for url in submitted_urls
    ]

    try:
        # Run blocking extraction off the event loop so other users are not
        # frozen while a batch is processing, and process up to 3 items at once.
        outcomes = await asyncio.to_thread(
            process_import_items,
            import_items,
            effective_api_key,
            ai_custom_columns,
        )
    except Exception:
        if credits_remaining is not None:
            refund_freemium_credits(current_user.id, chargeable_items_count)
        raise

    for outcome in outcomes:
        if outcome["ok"]:
            created_rows.append(outcome["row"])
        else:
            failures.append({"source": outcome["source"], "error": outcome["error"]})

    failed_chargeable = sum(1 for outcome in outcomes if not outcome["ok"] and outcome.get("chargeable"))
    credits_used = chargeable_items_count
    if credits_remaining is not None and failed_chargeable:
        refunded_balance = refund_freemium_credits(current_user.id, failed_chargeable)
        if refunded_balance is not None:
            credits_remaining = refunded_balance
        credits_used = chargeable_items_count - failed_chargeable

    persisted = []
    with session_scope() as db:
        default_template = project_brand_template(db, project_id)
        for source_type, source_name, source_url, stored_document, result in created_rows:
            product_path = save_extracted_asset(result.get("product_image_path"), PRODUCT_IMAGES_DIR, "products")
            diagram_path = save_extracted_asset(result.get("diagram_image_path"), DIAGRAM_IMAGES_DIR, "diagrams")
            polar_path = save_extracted_asset(result.get("polar_image_path"), POLAR_IMAGES_DIR, "polar")
            temp_dir = result.get("temp_dir")
            if temp_dir:
                shutil.rmtree(temp_dir, ignore_errors=True)
            metadata = normalize_import_metadata(result.get("metadata", {}))
            luminaire_reference_column = find_luminaire_reference_column_name(import_custom_columns)
            raw_custom_fields = result.get("custom_fields", {}) or {}
            normalized_custom_fields = {
                canonicalize_import_key(key): clean_optional_text(value)
                for key, value in raw_custom_fields.items()
            }
            custom_field_values = {}
            for column in import_custom_columns:
                column_name = column["name"]
                column_key = canonicalize_import_key(column_name)
                text_value = normalized_custom_fields.get(column_key)
                if text_value:
                    custom_field_values[column_name] = text_value
            if luminaire_reference_column:
                luminaire_reference = extract_luminaire_reference_from_source_name(source_name)
                if luminaire_reference:
                    custom_field_values[luminaire_reference_column] = luminaire_reference
            fill_empty_import_values(metadata, custom_field_values, import_custom_columns)
            row = Luminaire(
                project_id=project_id,
                source_type=source_type,
                source_name=source_name,
                source_url=source_url,
                document_path=stored_document,
                type_code=metadata.get("type"),
                manufacturer=metadata.get("manufacturer"),
                product_name=metadata.get("product"),
                spec_code=metadata.get("spec_code"),
                description=metadata.get("description"),
                colour=metadata.get("colour"),
                cct=metadata.get("cct"),
                cri=metadata.get("cri"),
                unit_watts=metadata.get("unit_watts"),
                dimming=metadata.get("dimming"),
                quantity=metadata.get("quantity"),
                ip_rating=metadata.get("ip_rating"),
                total_watts=metadata.get("total_watts"),
                pdf_template=(
                    metadata.get("template_pdf")
                    if str(metadata.get("template_pdf") or "").isdigit()
                    else str(default_template.id)
                ),
                diagram_image_path=diagram_path,
                product_image_path=product_path,
                polar_image_path=polar_path,
                custom_data=json.dumps(custom_field_values),
            )
            db.add(row)
            db.flush()
            persisted.append(serialize_luminaire(row))

        db.commit()

    return {
        "created": persisted,
        "failures": failures,
        "creditsUsed": credits_used if credits_remaining is not None else 0,
        "creditsRemaining": credits_remaining,
        "summary": {
            "requested": items_count,
            "created": len(persisted),
            "failed": len(failures),
        },
    }

# --- Settings & Designer API ---
@app.get("/api/settings/public")
def api_get_public_settings():
    """Returns non-sensitive settings needed by all pages (e.g. brand name)."""
    with session_scope() as db:
        return {
            "companyName": get_setting(db, "company_name", "LumenFlow"),
        }

@app.get("/api/settings")
def api_get_settings(current_user: User = Depends(get_admin_user)):
    with session_scope() as db:
        row = (
            db.query(AppSetting)
            .filter(AppSetting.key == "openai_api_key", AppSetting.project_id == None)  # noqa: E711
            .order_by(AppSetting.updated_at.desc())
            .first()
        )
        return {
            "companyName": get_setting(db, "company_name", "LumenFlow"),
            "items": [{
                "key": "openai_api_key",
                "label": "OpenAI API Key",
                "value": "Configured" if row and row.value else "Not set",
                "status": "Active" if row and row.value else "Missing",
                "updatedAt": row.updated_at.isoformat() if row and row.updated_at else "",
            }]
        }

@app.post("/api/settings")
async def api_save_settings(request: Request, current_user: User = Depends(get_admin_user)):
    data = await request.json()
    with session_scope() as db:
        key = data.get("openaiApiKey", "").strip()
        if key:
            set_setting(db, "openai_api_key", key)
        company_name = data.get("companyName", "").strip()
        if company_name:
            set_setting(db, "company_name", company_name)
        db.commit()
    return {"ok": True}


@app.get("/api/projects/{project_id}/specification-settings")
def api_get_specification_settings(project_id: int, current_user: User = Depends(get_current_user)):
    with session_scope() as db:
        project = get_project_or_404(db, project_id, current_user)
        return {
            "projectName": project.name,
            "logoUrl": project_logo_url(db, project_id),
            "locationCountry": get_setting(db, "specification_location_country", "", project_id=project_id),
            "revision": get_setting(db, "specification_revision", "A (Current)", project_id=project_id),
            "issueDate": get_setting(
                db,
                "specification_issue_date",
                datetime.utcnow().date().isoformat(),
                project_id=project_id,
            ),
        }


@app.put("/api/projects/{project_id}/specification-settings")
async def api_save_specification_settings(
    project_id: int,
    request: Request,
    current_user: User = Depends(get_current_user),
):
    payload = await request.json()
    with session_scope() as db:
        get_project_or_404(db, project_id, current_user)
        values = {
            "specification_location_country": str(payload.get("locationCountry", "") or "").strip(),
            "specification_revision": str(payload.get("revision", "") or "").strip() or "A (Current)",
            "specification_issue_date": str(payload.get("issueDate", "") or "").strip()
            or datetime.utcnow().date().isoformat(),
        }
        try:
            datetime.strptime(values["specification_issue_date"], "%Y-%m-%d")
        except ValueError:
            raise HTTPException(status_code=400, detail="Issue date must use YYYY-MM-DD format")
        for key, value in values.items():
            set_setting(db, key, value, project_id=project_id)
        db.commit()
        return {"ok": True}


@app.post("/api/projects/{project_id}/specification-logo")
async def api_upload_specification_logo(
    project_id: int,
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
):
    with session_scope() as db:
        get_project_or_404(db, project_id, current_user)

    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="Logo file is empty")
    if len(content) > 8 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Logo must be smaller than 8 MB")

    try:
        from PIL import Image

        with Image.open(BytesIO(content)) as source:
            source.verify()
        with Image.open(BytesIO(content)) as source:
            if source.width > 6000 or source.height > 6000:
                raise HTTPException(status_code=400, detail="Logo dimensions are too large")
            prepared = source.convert("RGBA")
            relative_path = f"branding/project-{project_id}-{uuid.uuid4().hex}.png"
            prepared.save(DATA_DIR / relative_path, "PNG", optimize=True)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=400, detail="Unsupported or invalid image file")

    with session_scope() as db:
        get_project_or_404(db, project_id, current_user)
        set_setting(db, "specification_logo_path", relative_path, project_id=project_id)
        db.commit()
        return {"logoUrl": to_public_file_url(relative_path)}


@app.get("/api/templates")
def api_list_templates(project_id: int, current_user: User = Depends(get_current_user)):
    with session_scope() as db:
        get_project_or_404(db, project_id, current_user)
        ensure_project_default_template(db, project_id)
        db.commit()
        rows = db.query(DesignTemplate).filter(DesignTemplate.project_id == project_id).order_by(DesignTemplate.updated_at.desc()).all()
        return {"templates": [serialize_template(r) for r in rows]}

@app.get("/api/templates/{template_id}")
def api_get_template(template_id: int, current_user: User = Depends(get_current_user)):
    with session_scope() as db:
        template = db.query(DesignTemplate).filter(DesignTemplate.id == template_id).first()
        if not template: raise HTTPException(404, "Template not found")
        get_project_or_404(db, template.project_id, current_user)
        if not template: raise HTTPException(404, "Template not found")
        return {"template": serialize_template(template)}

@app.post("/api/templates")
async def api_create_template(request: Request, project_id: int, current_user: User = Depends(get_current_user)):
    payload = await request.json()
    with session_scope() as db:
        get_project_or_404(db, project_id, current_user)
        first_lum = db.query(Luminaire).filter(Luminaire.project_id == project_id).first()
        template = DesignTemplate(
            project_id=project_id,
            luminaire_id=first_lum.id if first_lum else None,
            name=payload.get("name") or "Untitled",
            page_width=int(payload.get("pageWidth", 595)),
            page_height=int(payload.get("pageHeight", 842)),
            design_json=payload.get("designJson") or {"items": []}
        )
        db.add(template)
        db.commit()
        db.refresh(template)
        return {"template": serialize_template(template)}

@app.put("/api/templates/{template_id}")
async def api_update_template(template_id: int, project_id: int, request: Request, current_user: User = Depends(get_current_user)):
    payload = await request.json()
    with session_scope() as db:
        get_project_or_404(db, project_id, current_user)
        template = db.query(DesignTemplate).filter(DesignTemplate.id == template_id, DesignTemplate.project_id == project_id).first()
        if not template: raise HTTPException(404, "Template not found")
        template.name = payload.get("name") or template.name
        if "designJson" in payload:
            template.design_json = payload.get("designJson")
        template.updated_at = datetime.utcnow()
        db.commit()
        db.refresh(template)
        return {"template": serialize_template(template)}

@app.delete("/api/templates/{template_id}")
def api_delete_template(template_id: int, project_id: int, current_user: User = Depends(get_current_user)):
    with session_scope() as db:
        get_project_or_404(db, project_id, current_user)
        template = db.query(DesignTemplate).filter(DesignTemplate.id == template_id, DesignTemplate.project_id == project_id).first()
        if not template: raise HTTPException(404, "Template not found")
        if is_default_specification_template(template):
            raise HTTPException(400, "The project default template cannot be deleted. You can edit it in the designer.")
        db.delete(template)
        db.commit()
        return {"ok": True}

@app.post("/api/luminaires/{luminaire_id}/templates/{template_id}/export")
def api_export_template(luminaire_id: int, template_id: int, project_id: int, current_user: User = Depends(get_current_user)):
    with session_scope() as db:
        get_project_or_404(db, project_id, current_user)
        row = db.query(Luminaire).filter(Luminaire.id == luminaire_id, Luminaire.project_id == project_id).first()
        template = db.query(DesignTemplate).filter(DesignTemplate.id == template_id, DesignTemplate.project_id == project_id).first()
        if not row or not template: raise HTTPException(404, "Not found")
        rel = create_pdf_export(db, row, template)
        gen = GeneratedDocument(luminaire_id=luminaire_id, template_id=template_id, title=f"{template.name}.pdf", file_path=rel)
        db.add(gen)
        db.commit()
        db.refresh(gen)
        return {"document": serialize_generated_document(gen)}


def render_page_for_luminaire(pdf: canvas.Canvas, db: Session, row: Luminaire, template: DesignTemplate, ph: float):
    """Render one full page worth of items for a single luminaire row onto the canvas."""
    data = designer_context_from_luminaire(db, row)
    items = (template.design_json or {}).get("items", [])
    accent = normalize_accent(get_setting(db, "brand_accent", "", project_id=row.project_id)) if row.project_id else ""
    if accent and accent != DEFAULT_BRAND_ACCENT:
        items = [{**item, "style": branded_item_style(item.get("style") or {}, accent)} for item in items]
    for item in items:
        it = item.get("type")
        if it == "image" and item.get("bindingKey") == "projectLogoUrl" and not data.get("projectLogoUrl") and data.get("studioName"):
            render_text_item(pdf, {**item, "style": {**(item.get("style") or {}), "fontSize": 9, "fontWeight": "bold", "textAlign": "center", "verticalAlign": "middle"}}, data["studioName"], ph)
            continue
        if it in {"field", "text"}:
            val = data.get(item.get("bindingKey", ""), "") if it == "field" else item.get("text", "")
            if it == "field" and (val is None or str(val).strip() == ""):
                val = item.get("emptyText", "Not available")
            render_text_item(pdf, item, val, ph)
        elif it == "image":
            ip = get_internal_file_path(item.get("imageUrl") or data.get(item.get("bindingKey", ""), ""))
            render_image_item(pdf, item, ip, ph)
        elif it in {"shape", "line"}:
            x, y = float(item.get("x", 0)), ph - float(item.get("y", 0)) - float(item.get("height", 20))
            w, h = float(item.get("width", 100)), float(item.get("height", 20))
            style = item.get("style", {})
            pdf.setFillColor(hex_to_reportlab(style.get("backgroundColor"), colors.transparent))
            pdf.rect(x, y, w, h, stroke=0, fill=1)
            bc = hex_to_reportlab(style.get("borderColor"), colors.transparent)
            bw = float(style.get("borderWidth", 0))
            if bc != colors.transparent and bw > 0:
                pdf.setStrokeColor(bc)
                pdf.setLineWidth(bw)
                pdf.rect(x, y, w, h, stroke=1, fill=0)


@app.post("/api/bulk-export")
async def api_bulk_export(request: Request, project_id: int, current_user: User = Depends(get_current_user)):
    """Generate a single multi-page PDF — one page per luminaire — using a shared template."""
    body = await request.json()
    luminaire_ids: list[int] = body.get("luminaire_ids", [])
    template_id: int = body.get("template_id", 0)

    if not luminaire_ids or not template_id:
        raise HTTPException(400, "luminaire_ids and template_id are required")

    with session_scope() as db:
        get_project_or_404(db, project_id, current_user)
        template = db.query(DesignTemplate).filter(
            DesignTemplate.id == template_id,
            DesignTemplate.project_id == project_id
        ).first()
        if not template:
            raise HTTPException(404, "Template not found")

        pw, ph = page_dimensions(template.page_size, template.page_width, template.page_height)
        name = f"bulk_{uuid.uuid4().hex}.pdf"
        rel = f"exports/{name}"
        pdf = canvas.Canvas(str(EXPORTS_DIR / name), pagesize=(pw, ph))

        rendered = 0
        first_luminaire_id = luminaire_ids[0] if luminaire_ids else None

        for lid in luminaire_ids:
            row = db.query(Luminaire).filter(
                Luminaire.id == lid,
                Luminaire.project_id == project_id
            ).first()
            if not row:
                continue
            render_page_for_luminaire(pdf, db, row, template, ph)
            pdf.showPage()   # ← new page for each luminaire
            rendered += 1

        if rendered == 0:
            raise HTTPException(404, "No matching luminaires found")

        pdf.save()

        # Store a single GeneratedDocument entry for the bundle (linked to first luminaire)
        gen = GeneratedDocument(
            luminaire_id=first_luminaire_id,
            template_id=template_id,
            title=f"{template.name} ({rendered} pages).pdf",
            file_path=rel,
        )
        db.add(gen)
        db.commit()
        db.refresh(gen)
        return {"document": serialize_generated_document(gen), "pages": rendered}


@app.post("/api/luminaires/export-schedule")
async def api_export_schedule(
    payload: ScheduleExportRequest,
    project_id: int,
    current_user: User = Depends(get_current_user),
):
    luminaire_ids = dedupe_ints_preserve_order(payload.luminaire_ids)
    column_keys = normalize_column_keys(payload.column_keys)

    if not luminaire_ids:
        raise HTTPException(status_code=400, detail="At least one luminaire must be selected for export.")
    if not column_keys:
        raise HTTPException(status_code=400, detail="At least one visible column is required for export.")

    with session_scope() as db:
        project = get_project_or_404(db, project_id, current_user)

        rows = db.query(Luminaire).filter(
            Luminaire.project_id == project_id,
            Luminaire.id.in_(luminaire_ids),
        ).all()
        rows_by_id = {row.id: row for row in rows}
        ordered_rows = [rows_by_id[row_id] for row_id in luminaire_ids if row_id in rows_by_id]

        if not ordered_rows:
            raise HTTPException(status_code=404, detail="No matching luminaires found for this export.")

        custom_column_ids = []
        for key in column_keys:
            if not key.startswith("custom:"):
                continue
            try:
                custom_column_ids.append(int(key.split(":", 1)[1]))
            except Exception:
                continue

        custom_columns_by_id: dict[int, CustomColumn] = {}
        if custom_column_ids:
            custom_rows = db.query(CustomColumn).filter(
                CustomColumn.project_id == project_id,
                CustomColumn.id.in_(custom_column_ids),
            ).all()
            custom_columns_by_id = {row.id: row for row in custom_rows}

        template_ids = []
        for row in ordered_rows:
            if row.pdf_template and str(row.pdf_template).isdigit():
                template_ids.append(int(str(row.pdf_template)))

        template_names_by_id = {}
        if template_ids:
            template_rows = db.query(DesignTemplate).filter(
                DesignTemplate.project_id == project_id,
                DesignTemplate.id.in_(template_ids),
            ).all()
            template_names_by_id = {row.id: row.name for row in template_rows}

        export_rows = build_schedule_export_rows(
            ordered_rows,
            column_keys,
            payload.column_labels,
            custom_columns_by_id,
            template_names_by_id,
        )

        timestamp = datetime.utcnow().strftime("%Y%m%d-%H%M%S")
        base_name = f"{sanitize_export_filename_part(project.name)}-luminaire-schedule-{timestamp}"

        schedule_brand = get_schedule_brand(db, project_id)
        if payload.export_format == "excel":
            export_path = EXPORTS_DIR / f"{base_name}.xlsx"
            build_schedule_excel_export(export_path, project.name, export_rows, schedule_brand)
            media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        else:
            export_path = EXPORTS_DIR / f"{base_name}.pdf"
            build_schedule_pdf_export(export_path, project.name, export_rows, schedule_brand)
            media_type = "application/pdf"

        # Record the schedule so it appears in Documentation.
        try:
            log = json.loads(get_setting(db, "schedule_exports", "[]", project_id=project_id) or "[]")
        except Exception:
            log = []
        log.append({
            "file": f"exports/{export_path.name}",
            "format": "xlsx" if payload.export_format == "excel" else "pdf",
            "rows": len(export_rows),
            "revision": get_setting(db, "specification_revision", "A", project_id=project_id),
            "createdAt": datetime.utcnow().isoformat(),
        })
        set_setting(db, "schedule_exports", json.dumps(log[-200:]), project_id=project_id)
        db.commit()

    return FileResponse(
        path=str(export_path),
        media_type=media_type,
        filename=export_path.name,
    )



@app.get("/api/exports")
def api_list_exports(project_id: int, current_user: User = Depends(get_current_user)):
    with session_scope() as db:
        get_project_or_404(db, project_id, current_user)
        rows = db.query(GeneratedDocument).join(Luminaire).filter(Luminaire.project_id == project_id).order_by(GeneratedDocument.created_at.desc()).all()
        items = [{**serialize_generated_document(r), "type": "spec"} for r in rows]
        try:
            schedules = json.loads(get_setting(db, "schedule_exports", "[]", project_id=project_id) or "[]")
        except Exception:
            schedules = []
        for index, entry in enumerate(schedules):
            if not (DATA_DIR / str(entry.get("file", ""))).exists():
                continue
            items.append({
                "id": f"schedule-{index}",
                "type": "schedule",
                "title": f"Luminaire schedule ({entry.get('rows', 0)} row{'' if entry.get('rows', 0) == 1 else 's'}).{entry.get('format', 'xlsx')}",
                "fileUrl": to_public_file_url(entry.get("file")),
                "createdAt": entry.get("createdAt", ""),
                "revision": entry.get("revision", ""),
            })
        items.sort(key=lambda item: item.get("createdAt") or "", reverse=True)
        return {"items": items}

@app.post("/api/luminaires/{luminaire_id}/designer-assets")
async def api_upload_asset(luminaire_id: int, file: UploadFile = File(...), current_user: User = Depends(get_current_user)):
    with session_scope() as db:
        row = db.query(Luminaire).filter(Luminaire.id == luminaire_id).first()
        if not row:
            raise HTTPException(status_code=404, detail="Not found")
        get_project_or_404(db, row.project_id, current_user)
    suffix = Path(file.filename or ".png").suffix.lower()
    if suffix not in {".png", ".jpg", ".jpeg", ".webp"}:
        raise HTTPException(status_code=400, detail="Please upload a PNG, JPG or WebP image")
    content = await file.read()
    try:
        from PIL import Image

        with Image.open(BytesIO(content)) as source:
            source.verify()
    except Exception:
        raise HTTPException(status_code=400, detail="Unsupported or invalid image file")
    name = f"{luminaire_id}_{uuid.uuid4().hex}{suffix}"
    target = DESIGNER_IMAGES_DIR / name
    target.write_bytes(content)
    return {"imageUrl": to_public_file_url(f"images/designer/{name}"), "fileName": file.filename}


@app.post("/api/projects/{project_id}/designer-assets")
async def api_upload_project_asset(project_id: int, file: UploadFile = File(...), current_user: User = Depends(get_current_user)):
    with session_scope() as db:
        get_project_or_404(db, project_id, current_user)
    suffix = Path(file.filename or ".png").suffix.lower()
    if suffix not in {".png", ".jpg", ".jpeg", ".webp"}:
        raise HTTPException(status_code=400, detail="Please upload a PNG, JPG, or WebP image")
    name = f"project_{project_id}_{uuid.uuid4().hex}{suffix}"
    target = DESIGNER_IMAGES_DIR / name
    with target.open("wb") as output:
        shutil.copyfileobj(file.file, output)
    return {"imageUrl": to_public_file_url(f"images/designer/{name}"), "fileName": file.filename}


if __name__ == "__main__":
    reload_enabled = os.getenv("UVICORN_RELOAD", "").lower() in {"1", "true", "yes"}
    uvicorn.run("api:app", host="0.0.0.0", port=8000, reload=reload_enabled)
